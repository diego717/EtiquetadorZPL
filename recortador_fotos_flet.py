"""Edición visual de prueba del recortador, creada con Flet.

No reemplaza al recortador clásico: prueba una interfaz contemporánea y conserva
el procesamiento local de las fotografías.
"""

from __future__ import annotations

import asyncio
import base64
import csv
import io
import json
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import cv2
import flet as ft
import numpy as np
import qrcode
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageEnhance, ImageOps


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tif", ".tiff"}
APP_FOLDER = Path.home() / "RecortadorFotosFlet"
PROJECT_EXTENSION = ".recortador.json"


@dataclass
class FletPhoto:
    path: str
    crop: tuple[float, float, float, float]
    face_found: bool = False
    brightness: float = 1.0
    contrast: float = 1.0
    approval: str = "Pendiente"
    manual: bool = False
    issues: list[str] = field(default_factory=list)
    image: Image.Image | None = field(default=None, repr=False, compare=False)

    def project_data(self) -> dict[str, Any]:
        data = asdict(self)
        data.pop("image", None)
        return data


def to_base64(image: Image.Image, limit: tuple[int, int] = (1000, 1000)) -> str:
    image = image.convert("RGB").copy()
    image.thumbnail(limit, Image.Resampling.LANCZOS)
    output = io.BytesIO()
    image.save(output, "JPEG", quality=88)
    return base64.b64encode(output.getvalue()).decode("ascii")


def initial_crop(image: Image.Image, ratio: float = 3 / 4) -> tuple[float, float, float, float]:
    width, height = image.size
    if width / height > ratio:
        crop_height = height
        crop_width = crop_height * ratio
    else:
        crop_width = width
        crop_height = crop_width / ratio
    left, top = (width - crop_width) / 2, (height - crop_height) / 2
    return left, top, left + crop_width, top + crop_height


class FletRecortador:
    BG = "#101827"
    PANEL = "#172235"
    PANEL_ALT = "#202D42"
    TEXT = "#E8EEF8"
    MUTED = "#A6B4C8"
    TEAL = "#2DD4BF"
    TEAL_DARK = "#0F766E"
    GOLD = "#F6C453"
    DANGER = "#F87171"

    def __init__(self, page: ft.Page) -> None:
        self.page = page
        self.photos: list[FletPhoto] = []
        self.current: int | None = None
        self.output_dir: str = ""
        self.metadata: dict[str, dict[str, str]] = {}
        self.project_path: Path | None = None
        self.aspect_ratio = 3 / 4
        self.search_text = ""
        self.only_exceptions = False
        self._crop_reference: tuple[float, float, float, float] | None = None
        self._editor_scale = 1.0
        self._editor_origin = (0, 0)
        self._editor_size = (900, 600)
        self.file_picker = ft.FilePicker()
        self.page.services.append(self.file_picker)

        self.status = ft.Text("Cargá fotos para comenzar.", color=self.MUTED, size=12)
        self.summary = ft.Text("0 fotos · 0 pendientes", color=self.MUTED, size=12)
        self.photo_list = ft.ListView(expand=True, spacing=5, padding=8)
        self.preview = ft.Image(src="", width=self._editor_size[0], height=self._editor_size[1], fit=ft.BoxFit.CONTAIN, visible=False)
        self.preview_gesture = ft.GestureDetector(content=self.preview, mouse_cursor=ft.MouseCursor.MOVE, drag_interval=20, on_pan_update=self.move_crop_with_pointer)
        self.empty_preview = ft.Column(
            [ft.Icon(ft.Icons.IMAGE_SEARCH, size=56, color=self.MUTED), ft.Text("Seleccioná una foto para verla aquí", color=self.MUTED)],
            horizontal_alignment=ft.CrossAxisAlignment.CENTER,
            alignment=ft.MainAxisAlignment.CENTER,
            expand=True,
        )
        self.photo_name = ft.Text("Vista previa", size=18, weight=ft.FontWeight.W_600, color=self.TEXT)
        self.info = ft.Text("", size=12, color=self.MUTED)
        self.search = ft.TextField(hint_text="Buscar foto", prefix_icon=ft.Icons.SEARCH, dense=True, border_radius=10, on_change=self.on_search)
        self.approval = ft.Dropdown(
            label="Estado de aprobación",
            value="Pendiente",
            options=[ft.DropdownOption(key=value, text=value) for value in ("Pendiente", "Revisada", "Aprobada")],
            on_select=self.change_approval,
            dense=True,
        )
        self.brightness = ft.Slider(min=0.70, max=1.40, value=1.0, divisions=70, label="Brillo: {value}", active_color=self.TEAL, on_change=self.change_brightness)
        self.contrast = ft.Slider(min=0.75, max=1.45, value=1.0, divisions=70, label="Contraste: {value}", active_color=self.GOLD, on_change=self.change_contrast)
        self.zoom = ft.Slider(min=0.55, max=1.0, value=1.0, divisions=45, label="Área visible: {value}", active_color=self.TEAL, on_change=self.change_zoom)
        self.vertical = ft.Slider(min=-0.25, max=0.25, value=0, divisions=50, label="Desplazamiento vertical: {value}", active_color=self.GOLD, on_change=self.change_vertical)
        self.format = ft.Dropdown(label="Formato", value="JPG", options=[ft.DropdownOption(key=x, text=x) for x in ("JPG", "PNG", "WEBP")], dense=True)
        self.card_mode = ft.Switch(label="Generar tarjeta con QR", value=False, active_color=self.TEAL)
        self.brand_title = ft.TextField(label="Título de tarjeta", value="CREDENCIAL", dense=True)
        self.brand_color = ft.TextField(label="Color institucional", value="#0F766E", dense=True)
        self.output_label = ft.Text("Carpeta de salida sin seleccionar", color=self.MUTED, size=11)
        self.metadata_info = ft.Text("Sin datos CSV/Excel cargados.", color=self.MUTED, size=11)

    def show_snack(self, text: str, error: bool = False) -> None:
        self.status.value = text
        self.page.show_dialog(ft.SnackBar(ft.Text(text), bgcolor=self.DANGER if error else self.TEAL_DARK))
        self.page.update()

    @staticmethod
    def _key(path: str | Path) -> str:
        return Path(path).stem.casefold()

    def get_current(self) -> FletPhoto | None:
        if self.current is None or not 0 <= self.current < len(self.photos):
            return None
        return self.photos[self.current]

    def load_image(self, path: Path) -> Image.Image:
        with Image.open(path) as source:
            return ImageOps.exif_transpose(source).convert("RGB").copy()

    def crop_image(self, item: FletPhoto) -> Image.Image:
        image = item.image
        if image is None:
            image = self.load_image(Path(item.path))
            item.image = image
        image = ImageEnhance.Brightness(image).enhance(item.brightness)
        image = ImageEnhance.Contrast(image).enhance(item.contrast)
        return image.crop(tuple(round(value) for value in item.crop))

    def editor_canvas(self, item: FletPhoto) -> Image.Image:
        """Dibuja la foto original y el recuadro de encuadre arrastrable."""
        source = item.image
        if source is None:
            source = self.load_image(Path(item.path))
            item.image = source
        canvas_width, canvas_height = self._editor_size
        canvas = Image.new("RGB", (canvas_width, canvas_height), "#0B1220")
        scale = min((canvas_width - 28) / source.width, (canvas_height - 28) / source.height)
        draw_width, draw_height = round(source.width * scale), round(source.height * scale)
        origin_x, origin_y = (canvas_width - draw_width) // 2, (canvas_height - draw_height) // 2
        display = source.resize((draw_width, draw_height), Image.Resampling.LANCZOS)
        canvas.paste(display, (origin_x, origin_y))
        draw = ImageDraw.Draw(canvas)
        left, top, right, bottom = item.crop
        box = (origin_x + left * scale, origin_y + top * scale, origin_x + right * scale, origin_y + bottom * scale)
        draw.rectangle(box, outline="#FFFFFF", width=4)
        draw.rectangle((box[0] + 4, box[1] + 4, box[2] - 4, box[3] - 4), outline="#2DD4BF", width=2)
        for x, y in ((box[0], box[1]), (box[2], box[1]), (box[0], box[3]), (box[2], box[3])):
            draw.ellipse((x - 7, y - 7, x + 7, y + 7), fill="#2DD4BF", outline="#FFFFFF", width=1)
        self._editor_scale = scale
        self._editor_origin = (origin_x, origin_y)
        return canvas

    def detect_face_crop(self, item: FletPhoto) -> None:
        image = item.image or self.load_image(Path(item.path))
        item.image = image
        cascade = cv2.CascadeClassifier(str(Path(cv2.data.haarcascades) / "haarcascade_frontalface_default.xml"))
        gray = cv2.cvtColor(np.asarray(image), cv2.COLOR_RGB2GRAY)
        faces = cascade.detectMultiScale(gray, scaleFactor=1.12, minNeighbors=5, minSize=(50, 50))
        if len(faces) == 0:
            item.face_found = False
            item.crop = initial_crop(image, self.aspect_ratio)
            return
        x, y, width, height = max(faces, key=lambda face: face[2] * face[3])
        item.face_found = True
        crop_height = min(image.height, max(height * 3.8, image.height * 0.46))
        crop_width = crop_height * self.aspect_ratio
        if crop_width > image.width:
            crop_width = image.width
            crop_height = crop_width / self.aspect_ratio
        left = min(max(0, x + width / 2 - crop_width / 2), image.width - crop_width)
        top = min(max(0, y + height * 0.18 - crop_height * 0.36), image.height - crop_height)
        item.crop = left, top, left + crop_width, top + crop_height

    def analyze(self, item: FletPhoto) -> None:
        image = item.image or self.load_image(Path(item.path))
        item.image = image
        gray = cv2.cvtColor(np.asarray(image), cv2.COLOR_RGB2GRAY)
        issues = []
        if cv2.Laplacian(gray, cv2.CV_64F).var() < 45:
            issues.append("posible desenfoque")
        average = float(gray.mean())
        if average < 65:
            issues.append("muy oscura")
        elif average > 205:
            issues.append("muy clara")
        if not item.face_found:
            issues.append("rostro no detectado")
        item.issues = issues

    def update_preview(self) -> None:
        item = self.get_current()
        if item is None:
            self.preview.visible = False
            self.empty_preview.visible = True
            self.photo_name.value = "Vista previa"
            self.info.value = ""
            self.page.update()
            return
        result = self.crop_image(item)
        self.preview.src = "data:image/jpeg;base64," + to_base64(self.editor_canvas(item), self._editor_size)
        self.preview.visible = True
        self.empty_preview.visible = False
        self.photo_name.value = Path(item.path).name
        state = ", ".join(item.issues) if item.issues else "sin alertas"
        manual = " · ajustada manualmente" if item.manual else ""
        self.info.value = f"{result.width} × {result.height} px · {item.approval} · {state}{manual} · arrastrá el recuadro para moverlo"
        self.approval.value = item.approval
        self.brightness.value = item.brightness
        self.contrast.value = item.contrast
        self.page.update()

    def move_crop_with_pointer(self, event: ft.DragUpdateEvent) -> None:
        item = self.get_current()
        if item is None or item.image is None or event.local_delta is None:
            return
        dx, dy = event.local_delta.x / self._editor_scale, event.local_delta.y / self._editor_scale
        left, top, right, bottom = item.crop
        width, height = right - left, bottom - top
        left = min(max(0, left + dx), item.image.width - width)
        top = min(max(0, top + dy), item.image.height - height)
        item.crop = left, top, left + width, top + height
        item.manual = True
        self._crop_reference = item.crop
        self.zoom.value = 1.0
        self.vertical.value = 0
        self.update_preview()

    def rebuild_list(self) -> None:
        self.photo_list.controls.clear()
        visible = []
        for index, item in enumerate(self.photos):
            matches = not self.search_text or self.search_text in Path(item.path).name.casefold()
            exception = bool(item.issues) or item.approval != "Aprobada"
            if matches and (not self.only_exceptions or exception):
                visible.append((index, item))
        pending = sum(item.approval != "Aprobada" for item in self.photos)
        self.summary.value = f"{len(self.photos)} fotos · {pending} pendientes · {len(visible)} visibles"
        for index, item in visible:
            issue = bool(item.issues) or item.approval != "Aprobada"
            icon = ft.Icons.WARNING_AMBER_ROUNDED if issue else ft.Icons.CHECK_CIRCLE_ROUNDED
            color = self.GOLD if issue else self.TEAL
            selected = index == self.current
            self.photo_list.controls.append(
                ft.Container(
                    content=ft.ListTile(
                        leading=ft.Icon(icon, color=color),
                        title=ft.Text(Path(item.path).name, max_lines=1, overflow=ft.TextOverflow.ELLIPSIS),
                        subtitle=ft.Text(item.approval if not item.issues else ", ".join(item.issues[:2]), size=11),
                        on_click=lambda _event, choice=index: self.select_photo(choice),
                    ),
                    bgcolor=self.PANEL_ALT if selected else None,
                    border_radius=10,
                )
            )
        self.page.update()

    def select_photo(self, index: int) -> None:
        self.current = index
        item = self.get_current()
        self._crop_reference = item.crop if item else None
        self.zoom.value = 1.0
        self.vertical.value = 0
        self.rebuild_list()
        self.update_preview()

    def on_search(self, event: ft.ControlEvent) -> None:
        self.search_text = event.control.value.casefold().strip()
        self.rebuild_list()

    def toggle_exceptions(self, _event: ft.ControlEvent) -> None:
        self.only_exceptions = not self.only_exceptions
        self.rebuild_list()

    def change_right_menu(self, event: ft.ControlEvent) -> None:
        selected = event.control.selected
        if selected:
            self.show_right_menu(selected[0])

    def show_right_menu(self, section: str) -> None:
        if section == "editar":
            controls = [
                ft.Text("Ajustes de la foto", size=16, weight=ft.FontWeight.W_600), self.approval,
                ft.Divider(color="#334155"), ft.Text("Brillo", color=self.MUTED, size=12), self.brightness,
                ft.Text("Contraste", color=self.MUTED, size=12), self.contrast,
                ft.Text("Encuadre", color=self.MUTED, size=12), self.zoom, self.vertical,
                ft.Text("También podés arrastrar directamente el recuadro sobre la foto.", color=self.MUTED, size=11),
            ]
        elif section == "salida":
            controls = [
                ft.Text("Exportación", size=16, weight=ft.FontWeight.W_600), self.format,
                ft.OutlinedButton("Elegir carpeta de salida", icon=ft.Icons.DRIVE_FOLDER_UPLOAD, on_click=self.choose_output),
                self.output_label, ft.Divider(color="#334155"), self.card_mode,
                self.brand_title, self.brand_color,
                ft.FilledButton("Exportar lote", icon=ft.Icons.FILE_DOWNLOAD, on_click=self.export_all, style=ft.ButtonStyle(bgcolor="#B7791F", color="#101827")),
            ]
        elif section == "datos":
            controls = [
                ft.Text("Datos y QR", size=16, weight=ft.FontWeight.W_600),
                ft.OutlinedButton("Importar CSV / Excel", icon=ft.Icons.TABLE_VIEW, on_click=self.import_data),
                self.metadata_info,
                ft.Divider(color="#334155"),
                ft.Text("Para vincular, el archivo debe incluir una columna archivo. El QR toma qr, documento, doc o id.", color=self.MUTED, size=11),
            ]
        else:
            controls = [
                ft.Text("Control profesional", size=16, weight=ft.FontWeight.W_600),
                ft.OutlinedButton("Analizar lote", icon=ft.Icons.AUTO_AWESOME, on_click=self.analyze_all),
                ft.OutlinedButton("Sólo excepciones", icon=ft.Icons.FILTER_ALT, on_click=self.toggle_exceptions),
                ft.Divider(color="#334155"),
                ft.Text("Proyecto", size=13, weight=ft.FontWeight.W_600),
                ft.OutlinedButton("Guardar proyecto", icon=ft.Icons.SAVE, on_click=self.save_project),
                ft.OutlinedButton("Abrir proyecto", icon=ft.Icons.FOLDER_OPEN, on_click=self.load_project),
                ft.Text("Los estados Pendiente, Revisada y Aprobada se guardan dentro del proyecto.", color=self.MUTED, size=11),
            ]
        self.right_content.controls = controls
        self.page.update()

    def change_approval(self, event: ft.ControlEvent) -> None:
        item = self.get_current()
        if item:
            item.approval = event.control.value
            self.rebuild_list()
            self.update_preview()

    def change_brightness(self, event: ft.ControlEvent) -> None:
        item = self.get_current()
        if item:
            item.brightness = float(event.control.value)
            item.manual = True
            self.update_preview()

    def change_contrast(self, event: ft.ControlEvent) -> None:
        item = self.get_current()
        if item:
            item.contrast = float(event.control.value)
            item.manual = True
            self.update_preview()

    def transform_crop(self, zoom: float, vertical: float) -> None:
        item = self.get_current()
        if item is None or item.image is None:
            return
        left, top, right, bottom = self._crop_reference or item.crop
        width, height = right - left, bottom - top
        center_x, center_y = (left + right) / 2, (top + bottom) / 2 + vertical * item.image.height
        width, height = width * zoom, height * zoom
        left = min(max(0, center_x - width / 2), item.image.width - width)
        top = min(max(0, center_y - height / 2), item.image.height - height)
        item.crop = left, top, left + width, top + height
        item.manual = True
        self.update_preview()

    def change_zoom(self, event: ft.ControlEvent) -> None:
        self.transform_crop(float(event.control.value), float(self.vertical.value or 0))

    def change_vertical(self, event: ft.ControlEvent) -> None:
        self.transform_crop(float(self.zoom.value or 1), float(event.control.value))

    def auto_crop_current(self, _event: ft.ControlEvent) -> None:
        item = self.get_current()
        if item:
            self.detect_face_crop(item)
            self.analyze(item)
            item.manual = False
            self._crop_reference = item.crop
            self.zoom.value = 1.0
            self.vertical.value = 0
            self.rebuild_list()
            self.update_preview()

    def analyze_all(self, _event: ft.ControlEvent) -> None:
        for item in self.photos:
            if item.image is None:
                item.image = self.load_image(Path(item.path))
            if not item.face_found:
                self.detect_face_crop(item)
            self.analyze(item)
        self.rebuild_list()
        self.update_preview()
        self.show_snack("Análisis del lote terminado.")

    async def add_photos(self, _event: ft.ControlEvent) -> None:
        files = await self.file_picker.pick_files(
            dialog_title="Seleccionar fotografías",
            file_type=ft.FilePickerFileType.CUSTOM,
            allowed_extensions=[extension.lstrip(".") for extension in IMAGE_EXTENSIONS],
            allow_multiple=True,
        )
        if not files:
            return
        known = {str(Path(item.path).resolve()) for item in self.photos}
        added = 0
        for file in files:
            path = Path(file.path)
            if str(path.resolve()) in known:
                continue
            try:
                image = self.load_image(path)
            except OSError:
                continue
            self.photos.append(FletPhoto(path=str(path), crop=initial_crop(image), image=image))
            known.add(str(path.resolve()))
            added += 1
        if added and self.current is None:
            self.current = 0
        self.rebuild_list()
        self.update_preview()
        self.show_snack(f"Se cargaron {added} foto(s).")

    async def choose_output(self, _event: ft.ControlEvent) -> None:
        folder = await self.file_picker.get_directory_path(dialog_title="Elegir carpeta de salida")
        if folder:
            self.output_dir = folder
            self.output_label.value = Path(folder).name
            self.show_snack(f"Salida: {folder}")

    async def import_data(self, _event: ft.ControlEvent) -> None:
        files = await self.file_picker.pick_files(
            dialog_title="Importar CSV o Excel",
            file_type=ft.FilePickerFileType.CUSTOM,
            allowed_extensions=["csv", "xlsx", "xlsm"],
            allow_multiple=False,
        )
        if not files:
            return
        path = Path(files[0].path)
        rows: list[dict[str, str]] = []
        try:
            if path.suffix.casefold() == ".csv":
                with path.open(newline="", encoding="utf-8-sig") as file:
                    rows = [{str(key).casefold(): str(value or "") for key, value in row.items()} for row in csv.DictReader(file)]
            else:
                workbook = load_workbook(path, read_only=True, data_only=True)
                sheet = workbook.active
                values = list(sheet.iter_rows(values_only=True))
                headers = [str(value or "").strip().casefold() for value in values[0]]
                rows = [{headers[index]: str(value or "").strip() for index, value in enumerate(row)} for row in values[1:]]
                workbook.close()
        except (OSError, ValueError, IndexError) as error:
            self.show_snack(f"No se pudo importar: {error}", error=True)
            return
        self.metadata.clear()
        for row in rows:
            file_name = row.get("archivo", "") or row.get("file", "")
            if file_name:
                self.metadata[self._key(file_name)] = row
        self.metadata_info.value = f"{len(self.metadata)} registro(s) vinculados por la columna archivo."
        self.show_snack(f"Se importaron {len(self.metadata)} registro(s). Para QR se usa qr, documento, doc o id.")

    async def save_project(self, _event: ft.ControlEvent) -> None:
        APP_FOLDER.mkdir(parents=True, exist_ok=True)
        path = await self.file_picker.save_file(
            dialog_title="Guardar proyecto",
            file_name=(self.project_path.name if self.project_path else f"lote_{datetime.now():%Y%m%d}{PROJECT_EXTENSION}"),
            initial_directory=str(self.project_path.parent if self.project_path else APP_FOLDER),
        )
        if not path:
            return
        destination = Path(path)
        if not destination.name.endswith(PROJECT_EXTENSION):
            destination = destination.with_name(destination.name + PROJECT_EXTENSION)
        data = {
            "version": 1,
            "created": datetime.now().isoformat(timespec="seconds"),
            "output_dir": self.output_dir,
            "metadata": self.metadata,
            "photos": [item.project_data() for item in self.photos],
        }
        try:
            destination.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        except OSError as error:
            self.show_snack(f"No se pudo guardar: {error}", error=True)
            return
        self.project_path = destination
        self.show_snack(f"Proyecto guardado: {destination.name}")

    async def load_project(self, _event: ft.ControlEvent) -> None:
        files = await self.file_picker.pick_files(dialog_title="Abrir proyecto", file_type=ft.FilePickerFileType.CUSTOM, allowed_extensions=["json"], allow_multiple=False)
        if not files:
            return
        path = Path(files[0].path)
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            restored = []
            skipped = 0
            for saved in data.get("photos", []):
                source = Path(saved["path"])
                if not source.exists():
                    skipped += 1
                    continue
                image = self.load_image(source)
                crop = tuple(saved.get("crop", initial_crop(image)))
                restored.append(FletPhoto(
                    path=str(source), crop=crop, face_found=bool(saved.get("face_found", False)),
                    brightness=float(saved.get("brightness", 1)), contrast=float(saved.get("contrast", 1)),
                    approval=saved.get("approval", "Pendiente"), manual=bool(saved.get("manual", False)),
                    issues=list(saved.get("issues", [])), image=image,
                ))
        except (OSError, ValueError, KeyError) as error:
            self.show_snack(f"Proyecto inválido: {error}", error=True)
            return
        self.photos = restored
        self.metadata = data.get("metadata", {})
        self.output_dir = data.get("output_dir", "")
        self.output_label.value = Path(self.output_dir).name if self.output_dir else "Carpeta de salida sin seleccionar"
        self.metadata_info.value = f"{len(self.metadata)} registro(s) cargados desde el proyecto."
        self.project_path = path
        self.current = 0 if self.photos else None
        self.rebuild_list()
        self.update_preview()
        self.show_snack(f"Proyecto abierto: {len(restored)} fotos, {skipped} rutas no disponibles.")

    def qr_value(self, item: FletPhoto) -> str:
        row = self.metadata.get(self._key(item.path), {})
        return row.get("qr", "") or row.get("documento", "") or row.get("doc", "") or row.get("id", "") or Path(item.path).stem

    def card(self, photo: Image.Image, item: FletPhoto) -> Image.Image:
        row = self.metadata.get(self._key(item.path), {})
        title = self.brand_title.value.strip() or "CREDENCIAL"
        color = self.brand_color.value.strip()
        if not color.startswith("#") or len(color) != 7:
            color = "#0F766E"
        canvas = Image.new("RGB", (1016, 638), "white")
        draw = ImageDraw.Draw(canvas)
        draw.rectangle((0, 0, 1016, 96), fill=color)
        draw.text((38, 33), title.upper(), fill="white")
        portrait = photo.copy()
        portrait.thumbnail((330, 430), Image.Resampling.LANCZOS)
        canvas.paste(portrait, (44 + (330 - portrait.width) // 2, 138 + (430 - portrait.height) // 2))
        name = row.get("nombre", Path(item.path).stem)
        role = row.get("cargo", "")
        draw.text((430, 220), name.upper(), fill="#152238")
        draw.text((430, 260), role, fill="#52647A")
        qr = qrcode.make(self.qr_value(item)).convert("RGB")
        qr.thumbnail((145, 145), Image.Resampling.LANCZOS)
        canvas.paste(qr, (820, 455))
        return canvas

    async def export_all(self, _event: ft.ControlEvent) -> None:
        if not self.photos:
            self.show_snack("No hay fotos para exportar.", error=True)
            return
        if not self.output_dir:
            await self.choose_output(_event)
            if not self.output_dir:
                return
        output = Path(self.output_dir)
        try:
            output.mkdir(parents=True, exist_ok=True)
            extension, fmt = {"JPG": (".jpg", "JPEG"), "PNG": (".png", "PNG"), "WEBP": (".webp", "WEBP")}[self.format.value]
            records = []
            for item in self.photos:
                image = self.crop_image(item)
                if self.card_mode.value:
                    image = self.card(image, item)
                destination = output / f"{Path(item.path).stem}{extension}"
                counter = 2
                while destination.exists():
                    destination = output / f"{Path(item.path).stem}_{counter}{extension}"
                    counter += 1
                if fmt in ("JPEG", "WEBP"):
                    image.save(destination, fmt, quality=95)
                else:
                    image.save(destination, fmt)
                records.append((item, destination))
        except OSError as error:
            self.show_snack(f"Error de exportación: {error}", error=True)
            return
        log = output / f"registro_flet_{datetime.now():%Y%m%d_%H%M%S}.csv"
        with log.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file)
            writer.writerow(["Fecha", "Origen", "Exportado", "Aprobación", "Observaciones"])
            for item, destination in records:
                writer.writerow([datetime.now().isoformat(timespec="seconds"), item.path, destination.name, item.approval, "; ".join(item.issues)])
        self.show_snack(f"Exportadas {len(records)} fotos y creado {log.name}.")

    def build(self) -> None:
        self.page.title = "Recortador Studio — prueba Flet"
        self.page.bgcolor = self.BG
        self.page.padding = 0
        self.page.theme_mode = ft.ThemeMode.DARK
        self.page.theme = ft.Theme(color_scheme_seed=self.TEAL)
        self.page.window.width = 1440
        self.page.window.height = 900
        self.page.window.min_width = 1100
        self.page.window.min_height = 700

        header = ft.Container(
            content=ft.Row([
                ft.Column([ft.Text("RECORTADOR STUDIO", size=24, weight=ft.FontWeight.W_700, color=self.TEXT), ft.Text("Prueba visual Flet · flujo de aprobación · proyectos y QR", color=self.MUTED, size=12)], spacing=1),
                ft.Row([
                    ft.OutlinedButton("Abrir proyecto", icon=ft.Icons.FOLDER_OPEN, on_click=self.load_project),
                    ft.FilledButton("Guardar proyecto", icon=ft.Icons.SAVE, on_click=self.save_project, style=ft.ButtonStyle(bgcolor=self.TEAL_DARK)),
                ]),
            ], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
            padding=ft.Padding(22, 16, 22, 16), bgcolor="#0B1220",
        )
        toolbar = ft.Container(
            content=ft.Row([
                ft.FilledButton("Agregar fotos", icon=ft.Icons.ADD_PHOTO_ALTERNATE, on_click=self.add_photos, style=ft.ButtonStyle(bgcolor=self.TEAL_DARK)),
                ft.OutlinedButton("Detectar y analizar", icon=ft.Icons.AUTO_AWESOME, on_click=self.analyze_all),
                ft.OutlinedButton("Sólo excepciones", icon=ft.Icons.FILTER_ALT, on_click=self.toggle_exceptions),
            ], wrap=True, spacing=10),
            padding=ft.Padding(18, 12, 18, 12), bgcolor="#111C2C",
        )
        left = ft.Container(
            content=ft.Column([
                ft.Text("Lote de fotografías", size=16, weight=ft.FontWeight.W_600), self.summary,
                self.search,
                ft.OutlinedButton("Sólo excepciones", icon=ft.Icons.FILTER_ALT, on_click=self.toggle_exceptions),
                ft.Divider(color="#334155"), self.photo_list,
            ], expand=True), width=290, padding=12, bgcolor=self.PANEL, border_radius=14,
        )
        center = ft.Container(
            content=ft.Column([
                ft.Row([ft.Column([self.photo_name, self.info], spacing=2), ft.OutlinedButton("Detectar esta foto", icon=ft.Icons.FACE, on_click=self.auto_crop_current)], alignment=ft.MainAxisAlignment.SPACE_BETWEEN),
                ft.Divider(color="#334155"),
                ft.Container(content=ft.Stack([self.empty_preview, self.preview_gesture], expand=True, alignment=ft.Alignment.CENTER), expand=True, bgcolor="#0B1220", border_radius=14, padding=12),
            ], expand=True), expand=True, padding=14, bgcolor=self.PANEL, border_radius=14,
        )
        self.right_content = ft.Column([], scroll=ft.ScrollMode.AUTO, expand=True, spacing=10)
        menu = ft.SegmentedButton(
            segments=[
                ft.Segment(value="editar", icon=ft.Icons.TUNE, label="Editar"),
                ft.Segment(value="salida", icon=ft.Icons.FILE_DOWNLOAD, label="Salida"),
                ft.Segment(value="datos", icon=ft.Icons.TABLE_VIEW, label="Datos"),
                ft.Segment(value="pro", icon=ft.Icons.WORKSPACE_PREMIUM, label="Pro"),
            ],
            selected=["editar"], show_selected_icon=False, on_change=self.change_right_menu,
            style=ft.ButtonStyle(color={ft.ControlState.SELECTED: self.TEAL}),
        )
        right = ft.Container(
            content=ft.Column([
                ft.Text("Panel de trabajo", size=16, weight=ft.FontWeight.W_600), menu,
                ft.Divider(color="#334155"), self.right_content,
            ], expand=True), width=350, padding=14, bgcolor=self.PANEL, border_radius=14,
        )
        body = ft.Container(content=ft.Row([left, center, right], expand=True, spacing=12), padding=12, expand=True)
        footer = ft.Container(content=self.status, padding=ft.Padding(20, 9, 20, 9), bgcolor="#0B1220")
        self.page.add(ft.Column([header, toolbar, body, footer], expand=True, spacing=0))
        self.show_right_menu("editar")


async def main(page: ft.Page) -> None:
    app = FletRecortador(page)
    app.build()


if __name__ == "__main__":
    ft.run(main)
