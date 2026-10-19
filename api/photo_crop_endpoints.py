"""API local del recortador web de fotografias."""

from __future__ import annotations

import csv
import io
import json
import math
import shutil
import sys
import tempfile
import threading
import time
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Annotated, Literal
from urllib.parse import quote, urlsplit

from fastapi import APIRouter, HTTPException, Query, Request
from fastapi.responses import FileResponse, Response, StreamingResponse
from pydantic import BaseModel, ConfigDict, Field, field_validator, model_validator
from PIL import Image, ImageDraw, ImageFont, ImageOps, UnidentifiedImageError
import fitz
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
if str(PROJECT_ROOT / "src") not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT / "src"))

import face_detection
from card_renderer import (
    BARCODE_SYMBOLOGIES,
    BarcodeError,
    CardRenderer,
    FIELD_PATTERN,
    FontCache,
    barcode_svg,
    build_pdf,
    qr_svg,
    record_values,
    resolve_fields,
)
import card_magstripe
import photo_enhance
import card_printing
import card_vectors
from card_vector_pdf import NotVector, convert_to_paths, element_area, side_pdf
from system_fonts import public_font_catalog, resolve_font_variant
from photo_crop_service import (
    ExportTransform,
    SUPPORTED_EXTENSIONS,
    compute_face_crop,
    default_crop,
    encode_image,
    open_oriented_image,
    process_background,
    quality_issues,
    render_export,
    safe_output_stem,
)


MAX_PHOTOS = 250
MAX_FILE_BYTES = 30 * 1024 * 1024
MAX_SESSION_BYTES = 750 * 1024 * 1024
SESSION_TTL_SECONDS = 8 * 60 * 60
PREVIEW_LONG_SIDE = 1800
MAX_IMAGE_PIXELS = 80_000_000
MAX_PDF_BYTES = 100 * 1024 * 1024
MAX_PDF_PAGES = 60
MAX_DATA_BYTES = 15 * 1024 * 1024
MAX_PROJECT_BYTES = 850 * 1024 * 1024
MAX_ASSETS = 40
MAX_ASSET_BYTES = 20 * 1024 * 1024
MAX_CARD_ELEMENTS = 120
MAX_PATH_NODES = 40000
ASSET_EXTENSIONS = SUPPORTED_EXTENSIONS | {".svg"}


@dataclass
class PhotoRecord:
    id: str
    filename: str
    source_path: Path
    preview_path: Path
    width: int
    height: int
    size: int
    analysis: tuple | None = None
    # Fotos ya mejoradas o agrandadas, para no repetir el cálculo en cada tarjeta.
    enhance_cache: dict = field(default_factory=dict)

    def public(self, session_id: str) -> dict:
        return {
            "id": self.id,
            "filename": self.filename,
            "width": self.width,
            "height": self.height,
            "size": self.size,
            "preview_url": f"/api/photo-crop/sessions/{session_id}/photos/{self.id}/preview",
        }


@dataclass
class SessionAsset:
    id: str
    filename: str
    path: Path

    def public(self, session_id: str) -> dict:
        return {
            "asset_id": self.id,
            "filename": self.filename,
            "url": f"/api/photo-crop/sessions/{session_id}/assets/{self.id}",
        }


@dataclass
class PhotoSession:
    id: str
    folder: Path
    created_at: float = field(default_factory=time.time)
    touched_at: float = field(default_factory=time.time)
    photos: dict[str, PhotoRecord] = field(default_factory=dict)
    assets: dict[str, SessionAsset] = field(default_factory=dict)

    @property
    def total_bytes(self) -> int:
        return sum(photo.size for photo in self.photos.values())

    @property
    def logo_path(self) -> Path | None:
        asset = self.assets.get("logo")
        return asset.path if asset else None

    @property
    def card_background_path(self) -> Path | None:
        asset = self.assets.get("card_background")
        return asset.path if asset else None


class PhotoSessionStore:
    def __init__(self, root: Path | None = None) -> None:
        self.root = root or Path(tempfile.gettempdir()) / "EtiquetadorZPL" / "photo-crop"
        self.root.mkdir(parents=True, exist_ok=True)
        self.sessions: dict[str, PhotoSession] = {}
        self.lock = threading.RLock()
        self.cleanup()

    def cleanup(self) -> None:
        cutoff = time.time() - SESSION_TTL_SECONDS
        with self.lock:
            stale = [session_id for session_id, session in self.sessions.items() if session.touched_at < cutoff]
            for session_id in stale:
                self.delete(session_id)
            active = set(self.sessions)
        for folder in self.root.iterdir():
            if not folder.is_dir() or folder.name in active:
                continue
            try:
                if folder.stat().st_mtime < cutoff:
                    shutil.rmtree(folder, ignore_errors=True)
            except OSError:
                continue

    def create(self) -> PhotoSession:
        self.cleanup()
        session_id = uuid.uuid4().hex
        folder = self.root / session_id
        (folder / "source").mkdir(parents=True)
        (folder / "preview").mkdir()
        (folder / "assets").mkdir()
        session = PhotoSession(session_id, folder)
        with self.lock:
            self.sessions[session_id] = session
        return session

    def get(self, session_id: str) -> PhotoSession:
        with self.lock:
            session = self.sessions.get(session_id)
            if session is None:
                raise HTTPException(404, "La sesion de fotos no existe o vencio")
            session.touched_at = time.time()
            return session

    def delete(self, session_id: str) -> None:
        with self.lock:
            session = self.sessions.pop(session_id, None)
        if session:
            shutil.rmtree(session.folder, ignore_errors=True)


STORE = PhotoSessionStore()
DETECTION_LOCK = threading.Lock()
router = APIRouter(prefix="/api/photo-crop", tags=["photo-crop"])


def require_same_origin(request: Request) -> None:
    origin = request.headers.get("origin")
    if not origin:
        return
    parsed = urlsplit(origin)
    if parsed.netloc.casefold() != request.headers.get("host", "").casefold():
        raise HTTPException(403, "Origen no permitido para el recortador local")


def clean_filename(value: str) -> str:
    name = Path(value.replace("\\", "/")).name.strip()
    if not name or Path(name).suffix.casefold() not in SUPPORTED_EXTENSIONS:
        raise HTTPException(415, "Formato de imagen no admitido")
    return name[:220]


def save_preview(source_path: Path, preview_path: Path) -> tuple[int, int]:
    image = open_oriented_image(source_path)
    width, height = image.size
    if width * height > MAX_IMAGE_PIXELS:
        raise ValueError("La imagen tiene demasiados pixeles")
    preview = image.copy()
    preview.thumbnail((PREVIEW_LONG_SIDE, PREVIEW_LONG_SIDE), Image.Resampling.LANCZOS)
    preview.save(preview_path, "JPEG", quality=88, optimize=True)
    return width, height


async def read_limited_body(request: Request, limit: int, label: str) -> bytes:
    declared = request.headers.get("content-length")
    if declared:
        try:
            if int(declared) > limit:
                raise HTTPException(413, f"{label} supera el limite permitido")
        except ValueError:
            raise HTTPException(400, "Content-Length invalido")
    content = bytearray()
    async for chunk in request.stream():
        content.extend(chunk)
        if len(content) > limit:
            raise HTTPException(413, f"{label} supera el limite permitido")
    if not content:
        raise HTTPException(400, f"{label} esta vacio")
    return bytes(content)


def register_source_image(
    session: PhotoSession,
    source_path: Path,
    filename: str,
    *,
    photo_id: str | None = None,
) -> PhotoRecord:
    if len(session.photos) >= MAX_PHOTOS:
        source_path.unlink(missing_ok=True)
        raise HTTPException(413, f"El lote admite hasta {MAX_PHOTOS} fotos")
    size = source_path.stat().st_size
    if size > MAX_FILE_BYTES or session.total_bytes + size > MAX_SESSION_BYTES:
        source_path.unlink(missing_ok=True)
        raise HTTPException(413, "Se supero el limite de almacenamiento del lote")
    identifier = photo_id or uuid.uuid4().hex
    preview_path = session.folder / "preview" / f"{identifier}.jpg"
    try:
        width, height = save_preview(source_path, preview_path)
    except (UnidentifiedImageError, Image.DecompressionBombError, OSError, ValueError):
        source_path.unlink(missing_ok=True)
        preview_path.unlink(missing_ok=True)
        raise HTTPException(415, "El archivo no contiene una imagen valida")
    record = PhotoRecord(identifier, filename[:220], source_path, preview_path, width, height, size)
    session.photos[identifier] = record
    return record


def parse_tabular_file(filename: str, content: bytes) -> tuple[list[str], list[dict[str, str]]]:
    extension = Path(filename).suffix.casefold()
    if extension == ".csv":
        text = None
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise HTTPException(415, "No se pudo leer la codificacion del CSV")
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        headers = [str(value or "").strip() for value in (reader.fieldnames or [])]
        rows = [
            {header: str(row.get(header, "") or "").strip() for header in headers}
            for row in reader
        ]
    elif extension == ".xlsx":
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            first = next(iterator, None)
            headers = [str(value or "").strip() or f"Columna {index + 1}" for index, value in enumerate(first or [])]
            rows = []
            for values in iterator:
                row = {header: str(values[index] if index < len(values) and values[index] is not None else "").strip() for index, header in enumerate(headers)}
                if any(row.values()):
                    rows.append(row)
            workbook.close()
        except Exception as exc:
            raise HTTPException(415, "No se pudo leer la planilla Excel") from exc
    else:
        raise HTTPException(415, "Solo se admiten archivos CSV o XLSX")
    headers = headers[:100]
    rows = rows[:10_000]
    if not headers:
        raise HTTPException(400, "La planilla no contiene encabezados")
    return headers, rows


class DetectionRequest(BaseModel):
    photo_ids: list[str] | None = Field(default=None, max_length=MAX_PHOTOS)
    ratio: float = Field(default=0.75, gt=0.2, lt=5)
    hair_margin: float = Field(default=0.20, ge=0.08, le=0.35)
    eye_position: float = Field(default=40, ge=20, le=65)
    face_fill: float | None = Field(default=None, ge=0.25, le=0.90)


class ExportItem(BaseModel):
    photo_id: str
    crop: tuple[float, float, float, float]
    brightness: float = Field(default=1.0, ge=0.5, le=1.5)
    contrast: float = Field(default=1.0, ge=0.5, le=1.5)
    rotation: int = 0
    approved: bool = False
    issues: list[str] = Field(default_factory=list, max_length=20)
    metadata: dict[str, str] = Field(default_factory=dict)
    sequence: int | None = Field(default=None, ge=1, le=100_000)
    enhance: bool = True

    @field_validator("crop")
    @classmethod
    def validate_crop(cls, value: tuple[float, float, float, float]):
        if not all(math.isfinite(number) for number in value):
            raise ValueError("Coordenadas de recorte invalidas")
        if value[2] <= value[0] or value[3] <= value[1]:
            raise ValueError("El recorte debe tener ancho y alto positivos")
        return value

    @field_validator("rotation")
    @classmethod
    def validate_rotation(cls, value: int) -> int:
        normalized = value % 360
        if normalized not in {0, 90, 180, 270}:
            raise ValueError("Rotacion invalida")
        return normalized


HEX_COLOR = r"^#[0-9a-fA-F]{6}$"
ASSET_ID = r"^[a-z0-9_]{1,40}$"
FONT_ID = r"^(sans|serif|mono|system:[0-9a-f]{12})$"


class CardCondition(BaseModel):
    model_config = ConfigDict(extra="ignore")
    field: str = Field(default="", max_length=120)
    op: Literal["always", "equals", "not_equals", "contains", "empty", "not_empty"] = "always"
    value: str = Field(default="", max_length=200)


PathCoord = Annotated[float, Field(ge=-50, le=51)]


class PathNode(BaseModel):
    """Nodo de un trazado: punto y manejadores Bezier, en fracciones de la caja del objeto."""

    model_config = ConfigDict(extra="ignore")
    p: tuple[PathCoord, PathCoord]
    i: tuple[PathCoord, PathCoord] | None = None
    o: tuple[PathCoord, PathCoord] | None = None


class PathSubpath(BaseModel):
    model_config = ConfigDict(extra="ignore")
    closed: bool = False
    nodes: list[PathNode] = Field(min_length=1, max_length=MAX_PATH_NODES)


class CardGradient(BaseModel):
    model_config = ConfigDict(extra="ignore")
    kind: Literal["linear", "radial"] = "linear"
    start: str = Field(default="#ffffff", pattern=r"^#[0-9a-fA-F]{6}$")
    end: str = Field(default="#000000", pattern=r"^#[0-9a-fA-F]{6}$")
    angle: float = Field(default=0, ge=-360, le=360)


class CardElement(BaseModel):
    """Objeto libre de la credencial. Los campos que no aplican a un tipo se ignoran."""

    model_config = ConfigDict(extra="ignore")
    id: str = Field(max_length=40, pattern=r"^[A-Za-z0-9_-]+$")
    type: Literal["text", "photo", "image", "qr", "barcode", "rect", "ellipse", "line", "path"]
    name: str = Field(default="", max_length=80)
    x: float = Field(default=0, ge=-2, le=3)
    y: float = Field(default=0, ge=-2, le=3)
    w: float = Field(default=0.2, gt=0, le=4)
    h: float = Field(default=0.1, gt=0, le=4)
    rotation: float = Field(default=0, ge=-360, le=360)
    opacity: float = Field(default=1, ge=0, le=1)
    hidden: bool = False
    locked: bool = False
    condition: CardCondition | None = None
    text: str = Field(default="", max_length=1000)
    font: str = Field(default="sans", pattern=FONT_ID)
    size: float = Field(default=12, ge=2, le=300)
    bold: bool = False
    italic: bool = False
    underline: bool = False
    uppercase: bool = False
    color: str = Field(default="#17202b", pattern=HEX_COLOR)
    align: Literal["left", "center", "right"] = "left"
    valign: Literal["top", "middle", "bottom"] = "middle"
    fit: Literal["shrink", "wrap", "wrap_shrink", "none"] = "shrink"
    line_height: float = Field(default=1.15, ge=0.6, le=3)
    padding: float = Field(default=0, ge=0, le=50)
    fill: str | None = Field(default=None, pattern=HEX_COLOR)
    stroke_color: str | None = Field(default=None, pattern=HEX_COLOR)
    stroke_width: float = Field(default=0, ge=0, le=40)
    radius: float = Field(default=0, ge=0, le=200)
    shape: Literal["rect", "circle"] = "rect"
    image_fit: Literal["cover", "contain", "stretch"] = "cover"
    grayscale: bool = False
    asset: str | None = Field(default=None, pattern=ASSET_ID)
    value: str = Field(default="", max_length=1000)
    symbology: Literal["code128", "code39", "ean13", "ean8", "upca", "itf"] = "code128"
    show_text: bool = True
    ecc: Literal["L", "M", "Q", "H"] = "M"
    background: str | None = Field(default=None, pattern=HEX_COLOR)
    paths: list[PathSubpath] = Field(default_factory=list, max_length=4000)
    fill_rule: Literal["evenodd", "nonzero"] = "evenodd"
    gradient: CardGradient | None = None
    line_cap: Literal["butt", "round", "square"] = "butt"
    line_join: Literal["miter", "round", "bevel"] = "miter"
    dash: Literal["solid", "dash", "dot", "dashdot"] = "solid"
    recolor: str | None = Field(default=None, pattern=HEX_COLOR)

    @model_validator(mode="after")
    def limit_nodes(self) -> "CardElement":
        if sum(len(subpath.nodes) for subpath in self.paths) > MAX_PATH_NODES:
            raise ValueError(f"El trazado tiene más de {MAX_PATH_NODES} nodos")
        return self


class CardSide(BaseModel):
    model_config = ConfigDict(extra="ignore")
    background_color: str = Field(default="#ffffff", pattern=HEX_COLOR)
    background_image: str | None = Field(default=None, pattern=ASSET_ID)
    elements: list[CardElement] = Field(default_factory=list, max_length=MAX_CARD_ELEMENTS)


class MagStripe(BaseModel):
    """Datos de la banda magnética: plantillas con campos, por pista."""

    model_config = ConfigDict(extra="ignore")
    enabled: bool = False
    side: Literal["front", "back"] = "back"
    track1: str = Field(default="", max_length=200)
    track2: str = Field(default="", max_length=200)
    track3: str = Field(default="", max_length=200)


class CardDesign(BaseModel):
    model_config = ConfigDict(extra="ignore")
    version: int = 2
    width_mm: float = Field(default=85.6, ge=20, le=320)
    height_mm: float = Field(default=53.98, ge=20, le=320)
    double_sided: bool = False
    front: CardSide = Field(default_factory=CardSide)
    back: CardSide = Field(default_factory=CardSide)
    magstripe: MagStripe = Field(default_factory=MagStripe)


class PhotoOptions(BaseModel):
    background_mode: str = "original"
    background_strength: int = Field(default=55, ge=0, le=100)
    enhance: Literal["off", "soft", "normal", "strong"] = "off"
    upscale: bool = False

    @field_validator("background_mode")
    @classmethod
    def validate_background_mode(cls, value: str) -> str:
        normalized = value.casefold()
        if normalized not in {"original", "blur", "light"}:
            raise ValueError("Modo de fondo invalido")
        return normalized


class ExportRequest(PhotoOptions):
    items: list[ExportItem] = Field(min_length=1, max_length=MAX_PHOTOS)
    format: str = "JPG"
    width: int | None = Field(default=600, ge=64, le=6000)
    height: int | None = Field(default=800, ge=64, le=6000)
    suffix: str = Field(default="", max_length=60)
    card_design: CardDesign | None = None

    @field_validator("format")
    @classmethod
    def validate_format(cls, value: str) -> str:
        normalized = value.upper()
        if normalized not in {"JPG", "PNG", "WEBP"}:
            raise ValueError("Formato invalido")
        return normalized


class CardOutput(BaseModel):
    format: Literal["PDF", "PNG", "JPG"] = "PDF"
    dpi: int = Field(default=300, ge=72, le=600)
    sides: Literal["front", "back", "both"] = "front"
    pdf_layout: Literal["card", "sheet"] = "card"
    sheet: Literal["A4", "LETTER", "LEGAL"] = "A4"
    margin_mm: float = Field(default=10, ge=0, le=60)
    gap_mm: float = Field(default=4, ge=0, le=40)
    crop_marks: bool = True
    duplex_mirror: bool = True
    vector: bool = True
    filename_pattern: str = Field(default="{Archivo}", max_length=160)


class CardExportRequest(PhotoOptions):
    design: CardDesign
    items: list[ExportItem] = Field(min_length=1, max_length=MAX_PHOTOS)
    output: CardOutput = Field(default_factory=CardOutput)


class CardPreviewRequest(PhotoOptions):
    design: CardDesign
    item: ExportItem | None = None
    side: Literal["front", "back"] = "front"
    dpi: int = Field(default=150, ge=50, le=300)


class CardCheckRequest(BaseModel):
    design: CardDesign
    items: list[ExportItem] = Field(min_length=1, max_length=MAX_PHOTOS)
    sides: Literal["front", "back", "both"] = "front"
    upscale: bool = False


class EnhancePreviewRequest(PhotoOptions):
    item: ExportItem
    design: CardDesign | None = None
    width: int | None = Field(default=None, ge=64, le=6000)
    height: int | None = Field(default=None, ge=64, le=6000)


class CardPrintRequest(PhotoOptions):
    design: CardDesign
    items: list[ExportItem] = Field(min_length=1, max_length=MAX_PHOTOS)
    printer: str = Field(min_length=1, max_length=260)
    sides: Literal["front", "back", "both"] = "front"
    copies: int = Field(default=1, ge=1, le=10)
    rotate_back: bool = False
    encode: bool = False
    mag_start: list[str] = Field(default_factory=lambda: ["~1", "~2", "~3"], min_length=3, max_length=3)
    mag_end: list[str] = Field(default_factory=lambda: ["", "", ""], min_length=3, max_length=3)

    @field_validator("mag_start", "mag_end")
    @classmethod
    def validate_commands(cls, value: list[str]) -> list[str]:
        if any(len(command) > 12 or any(ord(char) < 32 for char in command) for command in value):
            raise ValueError("Comando de codificación inválido")
        return value


class ProjectRequest(BaseModel):
    state: dict = Field(default_factory=dict)


class VectorRequest(BaseModel):
    width_mm: float = Field(default=85.6, ge=20, le=320)
    height_mm: float = Field(default=53.98, ge=20, le=320)
    elements: list[CardElement] = Field(min_length=1, max_length=MAX_CARD_ELEMENTS)
    values: dict[str, str] = Field(default_factory=dict)


class BooleanRequest(VectorRequest):
    operation: Literal["union", "difference", "intersection", "xor"]


def render_portrait(
    record: PhotoRecord,
    item: ExportItem,
    options: PhotoOptions,
    output_size: tuple[int, int] | None,
    target: tuple[int, int] | None = None,
) -> Image.Image:
    """Recorte final de la foto. ``target`` es el tamaño en píxeles de la caja de la credencial."""
    level = options.enhance if item.enhance else "off"
    upscale = options.upscale and item.enhance
    transform = ExportTransform(item.crop, item.brightness, item.contrast, item.rotation)
    if level == "off" and not upscale:
        rendered = render_export(open_oriented_image(record.source_path), transform, output_size)
        return process_background(rendered, options.background_mode, options.background_strength)
    goal = output_size or target
    key = (tuple(round(value, 1) for value in item.crop), item.rotation, item.brightness, item.contrast, level, upscale, goal)
    cached = record.enhance_cache.get(key)
    if cached is None:
        rendered = render_export(open_oriented_image(record.source_path), transform, None)
        cached, _used_ai = photo_enhance.improve(rendered, goal, level=level, upscale=upscale)
        while len(record.enhance_cache) >= 3:
            record.enhance_cache.pop(next(iter(record.enhance_cache)))
        record.enhance_cache[key] = cached
    rendered = cached.copy()
    if output_size and rendered.size != output_size:
        rendered = rendered.resize(output_size, Image.Resampling.LANCZOS)
    return process_background(rendered, options.background_mode, options.background_strength)


def photo_target(design: CardDesign, sides: list[str], dpi: int) -> tuple[int, int] | None:
    """Píxeles de la caja de foto más grande del diseño a la resolución de salida."""
    sizes = [
        (element.w * design.width_mm / 25.4 * dpi, element.h * design.height_mm / 25.4 * dpi)
        for side in sides for element in getattr(design, side).elements if element.type == "photo" and not element.hidden
    ]
    if not sizes:
        return None
    return photo_enhance.ceil_size(max(size[0] for size in sizes), max(size[1] for size in sizes))


def load_assets(session: PhotoSession, design: CardDesign) -> dict[str, object]:
    """Abre una sola vez los recursos que usa el diseño."""
    wanted: set[str] = set()
    for side in (design.front, design.back):
        if side.background_image:
            wanted.add(side.background_image)
        wanted.update(element.asset for element in side.elements if element.type == "image" and element.asset)
    images: dict[str, object] = {}
    for asset_id in wanted:
        asset = session.assets.get(asset_id)
        if asset is None or not asset.path.exists():
            continue
        if asset.path.suffix.casefold() == ".svg":
            try:
                images[asset_id] = card_vectors.SvgAsset(asset.path.read_bytes())
            except Exception:  # noqa: BLE001 - SVG dañado: se informa como imagen sin cargar
                pass
            continue
        try:
            with Image.open(asset.path) as source:
                images[asset_id] = ImageOps.exif_transpose(source).convert("RGBA").copy()
        except (UnidentifiedImageError, OSError):
            continue
    return images


def card_renderer(session: PhotoSession, design: CardDesign, dpi: int) -> CardRenderer:
    fonts = FontCache(lambda family, bold, italic: resolve_font_variant(family, bold=bold, italic=italic))
    return CardRenderer(
        width_mm=design.width_mm,
        height_mm=design.height_mm,
        dpi=dpi,
        fonts=fonts,
        assets=load_assets(session, design),
    )


def magstripe_tracks(design: CardDesign, values: dict[str, str]) -> dict[int, str]:
    if not design.magstripe.enabled:
        return {}
    return card_magstripe.resolve_tracks(design.magstripe.model_dump(), values, resolve_fields)


def magstripe_issues(design: CardDesign, values: dict[str, str]) -> list[dict]:
    """Revisión de la banda: datos faltantes, caracteres no admitidos o demasiado largos."""
    if not design.magstripe.enabled:
        return []
    issues = []
    templates = design.magstripe.model_dump()
    tracks = magstripe_tracks(design, values)
    for track in card_magstripe.TRACKS:
        template = str(templates.get(f"track{track}") or "")
        if not template.strip():
            continue
        missing = [name.strip() for name in FIELD_PATTERN.findall(template) if not values.get(name.strip(), "").strip()]
        if missing:
            issues.append({"level": "warning", "element": "Banda magnética", "side": "",
                           "message": f"pista {track} sin dato en " + ", ".join(f"{{{name}}}" for name in dict.fromkeys(missing))})
        error = card_magstripe.track_error(track, tracks.get(track, ""))
        if error:
            issues.append({"level": "error", "element": "Banda magnética", "side": "", "message": error})
    if not any(tracks.values()):
        issues.append({"level": "error", "element": "Banda magnética", "side": "", "message": "no hay datos para grabar en ninguna pista"})
    return issues


def side_uses_photo(design: CardDesign, side: str) -> bool:
    return any(element.type == "photo" for element in getattr(design, side).elements)


def render_card_faces(
    renderer: CardRenderer,
    design: CardDesign,
    record: PhotoRecord,
    item: ExportItem,
    options: PhotoOptions,
    sides: list[str],
) -> list[Image.Image]:
    values = record_values(item.metadata, record.filename, item.sequence)
    needs_photo = any(side_uses_photo(design, side) for side in sides)
    portrait = render_portrait(record, item, options, None, photo_target(design, sides, renderer.dpi)) if needs_photo else None
    return [renderer.render_side(getattr(design, side).model_dump(), values, portrait) for side in sides]


def vector_card_faces(
    renderer: CardRenderer,
    design: CardDesign,
    record: PhotoRecord,
    item: ExportItem,
    options: PhotoOptions,
    sides: list[str],
) -> list:
    """Caras como documentos PDF vectoriales de una página."""
    values = record_values(item.metadata, record.filename, item.sequence)
    needs_photo = any(side_uses_photo(design, side) for side in sides)
    portrait = render_portrait(record, item, options, None, photo_target(design, sides, renderer.dpi)) if needs_photo else None
    return [side_pdf(renderer, getattr(design, side).model_dump(), values, portrait) for side in sides]


def render_item(
    session: PhotoSession,
    record: PhotoRecord,
    item: ExportItem,
    payload: ExportRequest,
    renderer: CardRenderer | None = None,
) -> Image.Image:
    if payload.card_design is not None:
        renderer = renderer or card_renderer(session, payload.card_design, 300)
        return render_card_faces(renderer, payload.card_design, record, item, payload, ["front"])[0]
    output_size = (payload.width, payload.height) if payload.width and payload.height else None
    return render_portrait(record, item, payload, output_size)


def attachment_headers(filename: str) -> dict[str, str]:
    """Content-Disposition compatible con nombres acentuados."""
    ascii_name = filename.encode("ascii", "replace").decode("ascii").replace("?", "_").replace('"', "_")
    return {"Content-Disposition": f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"}


@router.post("/sessions", status_code=201)
def create_session(request: Request) -> dict:
    require_same_origin(request)
    session = STORE.create()
    return {"session_id": session.id, "expires_in_seconds": SESSION_TTL_SECONDS}


@router.get("/fonts")
def list_system_fonts() -> dict:
    fonts = public_font_catalog()
    return {"fonts": fonts, "count": len(fonts)}


@router.get("/sessions/{session_id}")
def get_session(session_id: str) -> dict:
    session = STORE.get(session_id)
    return {"session_id": session.id, "photos": [photo.public(session.id) for photo in session.photos.values()]}


@router.post("/sessions/{session_id}/photos", status_code=201)
async def upload_photo(
    request: Request,
    session_id: str,
    filename: str = Query(..., min_length=1, max_length=500),
) -> dict:
    require_same_origin(request)
    session = STORE.get(session_id)
    name = clean_filename(filename)
    if len(session.photos) >= MAX_PHOTOS:
        raise HTTPException(413, f"El lote admite hasta {MAX_PHOTOS} fotos")
    declared = request.headers.get("content-length")
    if declared:
        try:
            if int(declared) > MAX_FILE_BYTES:
                raise HTTPException(413, "La foto supera el limite de 30 MB")
        except ValueError:
            raise HTTPException(400, "Content-Length invalido")

    photo_id = uuid.uuid4().hex
    extension = Path(name).suffix.casefold()
    source_path = session.folder / "source" / f"{photo_id}{extension}"
    preview_path = session.folder / "preview" / f"{photo_id}.jpg"
    size = 0
    try:
        with source_path.open("wb") as output:
            async for chunk in request.stream():
                size += len(chunk)
                if size > MAX_FILE_BYTES or session.total_bytes + size > MAX_SESSION_BYTES:
                    raise HTTPException(413, "Se supero el limite de almacenamiento del lote")
                output.write(chunk)
        if size == 0:
            raise HTTPException(400, "La foto esta vacia")
        width, height = save_preview(source_path, preview_path)
    except HTTPException:
        source_path.unlink(missing_ok=True)
        preview_path.unlink(missing_ok=True)
        raise
    except (UnidentifiedImageError, Image.DecompressionBombError, OSError, ValueError):
        source_path.unlink(missing_ok=True)
        preview_path.unlink(missing_ok=True)
        raise HTTPException(415, "El archivo no contiene una imagen valida")

    record = PhotoRecord(photo_id, name, source_path, preview_path, width, height, size)
    session.photos[photo_id] = record
    return record.public(session.id)


@router.post("/sessions/{session_id}/pdf", status_code=201)
async def import_pdf(
    request: Request,
    session_id: str,
    filename: str = Query(..., min_length=1, max_length=500),
) -> dict:
    require_same_origin(request)
    session = STORE.get(session_id)
    name = Path(filename.replace("\\", "/")).name.strip()[:220]
    if Path(name).suffix.casefold() != ".pdf":
        raise HTTPException(415, "El archivo debe ser un PDF")
    content = await read_limited_body(request, MAX_PDF_BYTES, "El PDF")
    pdf_path = session.folder / "assets" / f"import_{uuid.uuid4().hex}.pdf"
    pdf_path.write_bytes(content)
    created_ids: list[str] = []
    document = None
    try:
        document = fitz.open(pdf_path)
        if document.needs_pass:
            raise HTTPException(415, "El PDF esta protegido con contrasena")
        if document.page_count < 1:
            raise HTTPException(400, "El PDF no contiene paginas")
        if document.page_count > MAX_PDF_PAGES or len(session.photos) + document.page_count > MAX_PHOTOS:
            raise HTTPException(413, f"El PDF excede el limite de {MAX_PDF_PAGES} paginas o el lote de {MAX_PHOTOS} fotos")
        records: list[PhotoRecord] = []
        stem = safe_output_stem(name)
        for page_index in range(document.page_count):
            identifier = uuid.uuid4().hex
            source_path = session.folder / "source" / f"{identifier}.png"
            page = document.load_page(page_index)
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2.5, 2.5), alpha=False, colorspace=fitz.csRGB)
            pixmap.save(str(source_path))
            record = register_source_image(
                session,
                source_path,
                f"{stem}_pagina_{page_index + 1:02d}.png",
                photo_id=identifier,
            )
            created_ids.append(identifier)
            records.append(record)
        return {"photos": [record.public(session.id) for record in records], "pages": len(records)}
    except HTTPException:
        for identifier in created_ids:
            record = session.photos.pop(identifier, None)
            if record:
                record.source_path.unlink(missing_ok=True)
                record.preview_path.unlink(missing_ok=True)
        raise
    except (fitz.FileDataError, RuntimeError, OSError, ValueError) as exc:
        for identifier in created_ids:
            record = session.photos.pop(identifier, None)
            if record:
                record.source_path.unlink(missing_ok=True)
                record.preview_path.unlink(missing_ok=True)
        raise HTTPException(415, "No se pudo procesar el PDF") from exc
    finally:
        if document is not None:
            document.close()
        pdf_path.unlink(missing_ok=True)


@router.post("/sessions/{session_id}/data")
async def import_data(
    request: Request,
    session_id: str,
    filename: str = Query(..., min_length=1, max_length=500),
) -> dict:
    require_same_origin(request)
    STORE.get(session_id)
    name = Path(filename.replace("\\", "/")).name.strip()[:220]
    content = await read_limited_body(request, MAX_DATA_BYTES, "La planilla")
    headers, rows = parse_tabular_file(name, content)
    return {"filename": name, "headers": headers, "rows": rows, "count": len(rows)}


def validate_svg(content: bytes, label: str) -> None:
    if not card_vectors.is_svg(content):
        raise HTTPException(415, f"{label} no es un SVG válido")
    try:
        asset = card_vectors.SvgAsset(content)
        asset.render((32, 32))
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(415, f"{label} no es un SVG válido") from exc


def store_asset(session: PhotoSession, asset_id: str, filename: str, content: bytes, label: str, *, allow_svg: bool = False) -> SessionAsset:
    """Guarda una imagen de recursos validandola antes de exponerla."""
    extension = Path(filename).suffix.casefold()
    if extension == ".svg":
        if not allow_svg:
            raise HTTPException(415, f"{label} debe ser una imagen JPG, PNG, WEBP o BMP")
        validate_svg(content, label)
    path = session.folder / "assets" / f"{asset_id}{extension}"
    path.write_bytes(content)
    try:
        if extension != ".svg":
            with Image.open(path) as image:
                image.verify()
    except (UnidentifiedImageError, Image.DecompressionBombError, OSError):
        path.unlink(missing_ok=True)
        raise HTTPException(415, f"{label} no es una imagen valida")
    previous = session.assets.get(asset_id)
    if previous and previous.path != path:
        previous.path.unlink(missing_ok=True)
    asset = SessionAsset(asset_id, filename[:220], path)
    session.assets[asset_id] = asset
    return asset


def asset_response(session: PhotoSession, asset_id: str, missing: str) -> FileResponse:
    asset = session.assets.get(asset_id)
    if asset is None or not asset.path.exists():
        raise HTTPException(404, missing)
    headers = {"Cache-Control": "private, max-age=3600", "X-Content-Type-Options": "nosniff"}
    if asset.path.suffix.casefold() == ".svg":
        # Un SVG abierto directamente no debe poder ejecutar scripts en el origen local.
        headers["Content-Security-Policy"] = "default-src 'none'; style-src 'unsafe-inline'; img-src data:; sandbox"
        return FileResponse(asset.path, media_type="image/svg+xml", headers=headers)
    return FileResponse(asset.path, headers=headers)


@router.post("/sessions/{session_id}/assets", status_code=201)
async def upload_asset(
    request: Request,
    session_id: str,
    filename: str = Query(..., min_length=1, max_length=500),
) -> dict:
    require_same_origin(request)
    session = STORE.get(session_id)
    name = Path(filename.replace("\\", "/")).name.strip()[:220]
    if not name or Path(name).suffix.casefold() not in ASSET_EXTENSIONS:
        raise HTTPException(415, "Formato de imagen no admitido: usá JPG, PNG, WEBP, BMP o SVG")
    if len(session.assets) >= MAX_ASSETS:
        raise HTTPException(413, f"El diseño admite hasta {MAX_ASSETS} imagenes")
    content = await read_limited_body(request, MAX_ASSET_BYTES, "La imagen")
    asset = store_asset(session, f"a{uuid.uuid4().hex[:12]}", name, content, "La imagen", allow_svg=True)
    return asset.public(session.id)


@router.get("/sessions/{session_id}/assets/{asset_id}")
def get_asset(session_id: str, asset_id: str) -> FileResponse:
    return asset_response(STORE.get(session_id), asset_id, "Imagen no encontrada")


@router.post("/sessions/{session_id}/logo", status_code=201)
async def upload_logo(
    request: Request,
    session_id: str,
    filename: str = Query(..., min_length=1, max_length=500),
) -> dict:
    require_same_origin(request)
    session = STORE.get(session_id)
    name = clean_filename(filename)
    content = await read_limited_body(request, 10 * 1024 * 1024, "El logo")
    store_asset(session, "logo", name, content, "El logo")
    return {"filename": name, "url": f"/api/photo-crop/sessions/{session.id}/logo", "asset_id": "logo"}


@router.get("/sessions/{session_id}/logo")
def get_logo(session_id: str) -> FileResponse:
    return asset_response(STORE.get(session_id), "logo", "Logo no encontrado")


@router.post("/sessions/{session_id}/card-background", status_code=201)
async def upload_card_background(
    request: Request,
    session_id: str,
    filename: str = Query(..., min_length=1, max_length=500),
) -> dict:
    require_same_origin(request)
    session = STORE.get(session_id)
    name = clean_filename(filename)
    content = await read_limited_body(request, MAX_ASSET_BYTES, "El fondo")
    store_asset(session, "card_background", name, content, "El fondo")
    return {"filename": name, "url": f"/api/photo-crop/sessions/{session.id}/card-background", "asset_id": "card_background"}


@router.get("/sessions/{session_id}/card-background")
def get_card_background(session_id: str) -> FileResponse:
    return asset_response(STORE.get(session_id), "card_background", "Fondo no encontrado")


def optional_color(value: str | None) -> str | None:
    if not value:
        return None
    if len(value) != 7 or not value.startswith("#") or any(char not in "0123456789abcdefABCDEF" for char in value[1:]):
        raise HTTPException(422, "Color invalido")
    return value


SVG_HEADERS = {"Cache-Control": "private, max-age=600"}


def vector_renderer(payload: VectorRequest) -> CardRenderer:
    fonts = FontCache(lambda family, bold, italic: resolve_font_variant(family, bold=bold, italic=italic))
    return CardRenderer(width_mm=payload.width_mm, height_mm=payload.height_mm, dpi=600, fonts=fonts)


def vector_values(payload: VectorRequest) -> dict[str, str]:
    return record_values(payload.values, str(payload.values.get("Archivo") or "Nombre de ejemplo"))


@router.post("/vector/outline")
def vector_outline(payload: VectorRequest, request: Request) -> dict:
    """Convierte objetos a curvas editables (como Ctrl+Q en Corel)."""
    require_same_origin(request)
    renderer = vector_renderer(payload)
    values = vector_values(payload)
    results = []
    for element in payload.elements:
        try:
            converted = convert_to_paths(renderer, element.model_dump(), values)
        except NotVector as exc:
            raise HTTPException(422, f"{element.name or element.type}: no se puede convertir a curvas ({exc})") from exc
        results.append({"source": element.id, "elements": converted})
    return {"results": results}


@router.post("/vector/boolean")
def vector_boolean(payload: BooleanRequest, request: Request) -> dict:
    """Soldar, recortar, intersectar o excluir formas; devuelve el trazado resultante."""
    require_same_origin(request)
    renderer = vector_renderer(payload)
    values = vector_values(payload)
    card_mm = (payload.width_mm, payload.height_mm)
    geometries = []
    for element in payload.elements:
        if element.type in {"photo", "image"}:
            raise HTTPException(422, "Las fotos e imágenes no se pueden combinar: usá formas, trazados o textos")
        try:
            geometries.append(element_area(renderer, element.model_dump(), values, card_mm))
        except NotVector as exc:
            raise HTTPException(422, f"{element.name or element.type}: {exc}") from exc
    try:
        result = card_vectors.boolean_geometry(payload.operation, geometries)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    fitted = card_vectors.fit_paths_to_box(card_vectors.geometry_paths(result, simplify=0.004))
    if fitted is None:
        raise HTTPException(422, "El resultado quedó vacío: las formas no se superponen")
    unit, (left, top, width, height) = fitted
    total_nodes = sum(len(subpath["nodes"]) for subpath in unit)
    if total_nodes > MAX_PATH_NODES:
        raise HTTPException(422, "El resultado es demasiado complejo")
    return {"element": {
        "x": left / payload.width_mm, "y": top / payload.height_mm,
        "w": width / payload.width_mm, "h": height / payload.height_mm,
        "rotation": 0, "fill_rule": "evenodd", "paths": card_vectors.serialize_paths(unit),
    }}


@router.get("/qr.svg")
def qr_preview(
    value: str = Query(..., min_length=1, max_length=1000),
    ecc: Literal["L", "M", "Q", "H"] = "M",
    color: str = "#000000",
    background: str | None = None,
) -> Response:
    svg = qr_svg(value, ecc, optional_color(color) or "#000000", optional_color(background))
    return Response(svg, media_type="image/svg+xml", headers=SVG_HEADERS)


@router.get("/barcode.svg")
def barcode_preview(
    value: str = Query(..., min_length=1, max_length=200),
    symbology: str = "code128",
    show_text: bool = True,
    color: str = "#000000",
    background: str | None = None,
) -> Response:
    if symbology not in BARCODE_SYMBOLOGIES:
        raise HTTPException(422, "Simbologia no admitida")
    try:
        svg = barcode_svg(symbology, value, show_text, optional_color(color) or "#000000", optional_color(background))
    except BarcodeError as exc:
        raise HTTPException(422, f"Codigo invalido: {exc}") from exc
    return Response(svg, media_type="image/svg+xml", headers=SVG_HEADERS)


@router.get("/sessions/{session_id}/photos/{photo_id}/preview")
def photo_preview(session_id: str, photo_id: str) -> FileResponse:
    photo = STORE.get(session_id).photos.get(photo_id)
    if photo is None:
        raise HTTPException(404, "Foto no encontrada")
    return FileResponse(photo.preview_path, media_type="image/jpeg", headers={"Cache-Control": "private, max-age=3600"})


@router.post("/sessions/{session_id}/detect")
def detect_photos(session_id: str, payload: DetectionRequest, request: Request) -> dict:
    require_same_origin(request)
    session = STORE.get(session_id)
    selected_ids = payload.photo_ids or list(session.photos)
    unknown = [photo_id for photo_id in selected_ids if photo_id not in session.photos]
    if unknown:
        raise HTTPException(404, "Una o mas fotos no pertenecen a la sesion")
    results = []
    for photo_id in selected_ids:
        record = session.photos[photo_id]
        if record.analysis is None:
            # La detección es lo costoso: se guarda para poder reencuadrar el lote al instante.
            image = open_oriented_image(record.source_path)
            with DETECTION_LOCK:
                face = face_detection.detect_primary_face(image)
            record.analysis = (image.size, face, quality_issues(image, face.box if face else None))
        image_size, face, issues = record.analysis
        if face:
            crop = compute_face_crop(
                image_size, face.box, face.eye_line, payload.hair_margin,
                payload.ratio, payload.eye_position, payload.face_fill,
            )
            box = list(face.box)
            eye_line = face.eye_line
            eyes = [list(point) for point in face.eye_points] if face.eye_points else None
            engine = face.engine
            score = face.score
        else:
            crop = default_crop(image_size, payload.ratio)
            box = eye_line = eyes = None
            engine = face_detection.engine()
            score = 0.0
        results.append({
            "photo_id": photo_id,
            "face_found": face is not None,
            "face_box": box,
            "eye_line": eye_line,
            "eye_points": eyes,
            "engine": engine,
            "score": score,
            "crop": list(crop),
            "issues": list(issues),
        })
    return {"results": results, "engine_note": face_detection.engine_note()}


def unique_name(base: str, extension: str, used: set[str]) -> str:
    candidate = f"{base}{extension}"
    counter = 2
    while candidate.casefold() in used:
        candidate = f"{base}_{counter}{extension}"
        counter += 1
    used.add(candidate.casefold())
    return candidate


def session_record(session: PhotoSession, photo_id: str) -> PhotoRecord:
    record = session.photos.get(photo_id)
    if record is None:
        raise HTTPException(404, "Una foto del lote ya no esta disponible")
    return record


@router.post("/sessions/{session_id}/export")
def export_photos(session_id: str, payload: ExportRequest, request: Request) -> StreamingResponse:
    require_same_origin(request)
    session = STORE.get(session_id)
    if len({item.photo_id for item in payload.items}) != len(payload.items):
        raise HTTPException(400, "El lote contiene fotos repetidas")
    renderer = card_renderer(session, payload.card_design, 300) if payload.card_design else None
    archive = io.BytesIO()
    used_names: set[str] = set()
    log_rows: list[list[str]] = []
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as bundle:
        for item in payload.items:
            record = session_record(session, item.photo_id)
            rendered = render_item(session, record, item, payload, renderer)
            encoded, extension, _mime = encode_image(rendered, payload.format)
            candidate = unique_name(safe_output_stem(record.filename, payload.suffix), extension, used_names)
            bundle.writestr(candidate, encoded)
            status = "Revisada" if item.approved else (item.issues[0] if item.issues else "Exportada")
            log_rows.append([datetime.now().isoformat(timespec="seconds"), record.filename, candidate, status])
        log_buffer = io.StringIO(newline="")
        writer = csv.writer(log_buffer)
        writer.writerow(["Fecha y hora", "Archivo origen", "Archivo exportado", "Estado"])
        writer.writerows(log_rows)
        bundle.writestr("registro_lote.csv", "﻿" + log_buffer.getvalue())
    archive.seek(0)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return StreamingResponse(archive, media_type="application/zip", headers=attachment_headers(f"fotos_recortadas_{stamp}.zip"))


@router.post("/sessions/{session_id}/review-pdf")
def review_pdf(session_id: str, payload: ExportRequest, request: Request) -> StreamingResponse:
    require_same_origin(request)
    session = STORE.get(session_id)
    renderer = card_renderer(session, payload.card_design, 300) if payload.card_design else None
    pages: list[Image.Image] = []
    page = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(page)
    title_font = ImageFont.load_default(size=28)
    text_font = ImageFont.load_default(size=18)
    draw.text((55, 35), "Hoja de revision de fotografias", fill=(22, 32, 45), font=title_font)
    cell_width, cell_height = 376, 510
    for index, item in enumerate(payload.items):
        record = session_record(session, item.photo_id)
        slot = index % 9
        if slot == 0 and index:
            pages.append(page)
            page = Image.new("RGB", (1240, 1754), "white")
            draw = ImageDraw.Draw(page)
            draw.text((55, 35), "Hoja de revision de fotografias", fill=(22, 32, 45), font=title_font)
        column, row = slot % 3, slot // 3
        x, y = 45 + column * 398, 95 + row * 540
        rendered = render_item(session, record, item, payload, renderer)
        preview = ImageOps.contain(rendered.convert("RGB"), (cell_width - 24, 390), Image.Resampling.LANCZOS)
        page.paste(preview, (x + (cell_width - preview.width) // 2, y + 10))
        draw.rectangle((x, y, x + cell_width, y + cell_height), outline=(215, 222, 231), width=2)
        label = Path(record.filename).stem[:38]
        status = "Revisada" if item.approved else (item.issues[0][:42] if item.issues else "Lista")
        draw.text((x + 12, y + 414), label, fill=(22, 32, 45), font=text_font)
        draw.text((x + 12, y + 452), status, fill=(29, 126, 112) if item.approved else (177, 112, 31), font=text_font)
    pages.append(page)
    output = io.BytesIO()
    pages[0].save(output, "PDF", resolution=150, save_all=True, append_images=pages[1:])
    output.seek(0)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return StreamingResponse(output, media_type="application/pdf", headers=attachment_headers(f"revision_fotos_{stamp}.pdf"))


@router.post("/sessions/{session_id}/cards/preview")
def card_preview(session_id: str, payload: CardPreviewRequest, request: Request) -> Response:
    """Render exacto de una cara, el mismo que se usa al exportar."""
    require_same_origin(request)
    session = STORE.get(session_id)
    renderer = card_renderer(session, payload.design, payload.dpi)
    side = getattr(payload.design, payload.side).model_dump()
    if payload.item is not None:
        record = session_record(session, payload.item.photo_id)
        image = render_card_faces(renderer, payload.design, record, payload.item, payload, [payload.side])[0]
    else:
        image = renderer.render_side(side, record_values({}, "Nombre de ejemplo"), None)
    output = io.BytesIO()
    image.save(output, "PNG", optimize=False)
    return Response(output.getvalue(), media_type="image/png", headers={"Cache-Control": "no-store"})


@router.post("/sessions/{session_id}/cards/export")
def export_cards(session_id: str, payload: CardExportRequest, request: Request) -> StreamingResponse:
    """Exporta credenciales individuales o por lote a PDF, PNG o JPG."""
    require_same_origin(request)
    session = STORE.get(session_id)
    output = payload.output
    design = payload.design
    if output.sides == "both" and not design.double_sided:
        raise HTTPException(400, "El diseño no tiene dorso habilitado")
    if output.sides == "back" and not design.double_sided:
        raise HTTPException(400, "El diseño no tiene dorso habilitado")
    sides = ["front", "back"] if output.sides == "both" else [output.sides]
    if output.dpi > 300 and len(payload.items) * len(sides) > 120:
        raise HTTPException(413, "A 600 dpi se pueden exportar hasta 120 caras por vez")
    renderer = card_renderer(session, design, output.dpi)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    vector_pdf = output.format == "PDF" and output.vector
    rendered: list[tuple[PhotoRecord, ExportItem, list]] = []
    for item in payload.items:
        record = session_record(session, item.photo_id)
        if vector_pdf:
            rendered.append((record, item, vector_card_faces(renderer, design, record, item, payload, sides)))
        else:
            rendered.append((record, item, render_card_faces(renderer, design, record, item, payload, sides)))

    if output.format == "PDF":
        try:
            pdf = build_pdf(
                [faces for _record, _item, faces in rendered],
                (design.width_mm, design.height_mm),
                layout=output.pdf_layout,
                sheet=output.sheet,
                margin_mm=output.margin_mm,
                gap_mm=output.gap_mm,
                crop_marks=output.crop_marks,
                duplex_mirror=output.duplex_mirror,
            )
        except ValueError as exc:
            raise HTTPException(400, str(exc)) from exc
        if len(rendered) == 1:
            record, item, _faces = rendered[0]
            values = record_values(item.metadata, record.filename, item.sequence)
            name = safe_output_stem(resolve_fields(output.filename_pattern, values) or record.filename)
        else:
            name = f"credenciales_{stamp}"
        return StreamingResponse(io.BytesIO(pdf), media_type="application/pdf", headers=attachment_headers(f"{name}.pdf"))

    files: list[tuple[str, bytes]] = []
    used_names: set[str] = set()
    side_suffix = {"front": "_frente", "back": "_dorso"}
    for record, item, faces in rendered:
        values = record_values(item.metadata, record.filename, item.sequence)
        base = safe_output_stem(resolve_fields(output.filename_pattern, values).strip() or record.filename)
        for side, face in zip(sides, faces):
            buffer = io.BytesIO()
            if output.format == "PNG":
                face.save(buffer, "PNG", dpi=(output.dpi, output.dpi), optimize=True)
                extension = ".png"
            else:
                face.save(buffer, "JPEG", quality=95, subsampling=0, dpi=(output.dpi, output.dpi))
                extension = ".jpg"
            suffix = side_suffix[side] if len(sides) > 1 else ""
            files.append((unique_name(f"{base}{suffix}", extension, used_names), buffer.getvalue()))

    if len(files) == 1:
        name, content = files[0]
        media = "image/png" if output.format == "PNG" else "image/jpeg"
        return StreamingResponse(io.BytesIO(content), media_type=media, headers=attachment_headers(name))
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_STORED) as bundle:
        for name, content in files:
            bundle.writestr(name, content)
    archive.seek(0)
    return StreamingResponse(archive, media_type="application/zip", headers=attachment_headers(f"credenciales_{stamp}.zip"))


def requested_sides(design: CardDesign, sides: str) -> list[str]:
    if sides != "front" and not design.double_sided:
        raise HTTPException(400, "El diseño no tiene dorso habilitado")
    return ["front", "back"] if sides == "both" else [sides]


@router.post("/sessions/{session_id}/cards/check")
def check_cards(session_id: str, payload: CardCheckRequest, request: Request) -> dict:
    """Revisión previa a exportar o imprimir; no compone imágenes."""
    require_same_origin(request)
    session = STORE.get(session_id)
    sides = requested_sides(payload.design, payload.sides)
    renderer = card_renderer(session, payload.design, 300)
    side_labels = {"front": "Frente", "back": "Dorso"}
    results = []
    for item in payload.items:
        record = session_record(session, item.photo_id)
        values = record_values(item.metadata, record.filename, item.sequence)
        issues = []
        for side in sides:
            side_data = getattr(payload.design, side)
            for issue in renderer.inspect_side(side_data.model_dump(), values):
                issues.append({**issue, "side": side_labels[side]})
        issues.extend(magstripe_issues(payload.design, values))
        target = photo_target(payload.design, sides, 300)
        crop_size = (round(item.crop[2] - item.crop[0]), round(item.crop[3] - item.crop[1]))
        if target and photo_enhance.needs_upscale(crop_size, target):
            rescue = payload.upscale and item.enhance and photo_enhance.model_path() is not None
            too_small = photo_enhance.missing_scale(crop_size, target) > photo_enhance.TOO_SMALL
            issues.append({
                "level": "warning", "element": "Foto", "side": "",
                "message": f"foto de baja resolución ({crop_size[0]}×{crop_size[1]} px; la caja pide {target[0]}×{target[1]} a 300 dpi)"
                + (" · conviene pedir otra foto" if too_small else " · se agranda con IA" if rescue else " · puede verse pixelada"),
            })
        if not item.approved and item.issues:
            issues.append({"level": "warning", "element": "Foto", "side": "", "message": item.issues[0].lower()})
        if issues:
            results.append({"photo_id": item.photo_id, "filename": record.filename, "issues": issues})
    errors = sum(1 for result in results for issue in result["issues"] if issue["level"] == "error")
    return {"checked": len(payload.items), "with_issues": len(results), "errors": errors, "results": results}


@router.post("/sessions/{session_id}/enhance-preview")
def enhance_preview(session_id: str, payload: EnhancePreviewRequest, request: Request) -> Response:
    """Comparación antes / después de la mejora de imagen para una foto."""
    require_same_origin(request)
    session = STORE.get(session_id)
    record = session_record(session, payload.item.photo_id)
    output_size = (payload.width, payload.height) if payload.width and payload.height else None
    target = photo_target(payload.design, ["front", "back"] if payload.design.double_sided else ["front"], 300) if payload.design else None
    goal = output_size or target
    plain_options = payload.model_copy(update={"enhance": "off", "upscale": False})
    before = render_portrait(record, payload.item, plain_options, output_size)
    after = render_portrait(record, payload.item.model_copy(update={"enhance": True}), payload, output_size, target)
    used_ai = bool(goal) and payload.upscale and photo_enhance.needs_upscale(before.size, goal) and photo_enhance.super_resolver() is not None
    height = 560
    panels = [image.resize((max(1, round(image.width * height / image.height)), height), Image.Resampling.LANCZOS) for image in (before, after)]
    gap = 12
    canvas = Image.new("RGB", (panels[0].width + panels[1].width + gap, height + 34), (12, 20, 32))
    canvas.paste(panels[0], (0, 34))
    canvas.paste(panels[1], (panels[0].width + gap, 34))
    draw = ImageDraw.Draw(canvas)
    try:
        font = ImageFont.truetype("arialbd.ttf", 18)
    except OSError:
        font = ImageFont.load_default()
    draw.text((8, 8), f"Antes · {before.width}×{before.height} px", font=font, fill=(200, 210, 222))
    draw.text((panels[0].width + gap + 8, 8), "Después" + (" · agrandada con IA" if used_ai else ""), font=font, fill=(120, 230, 210))
    output = io.BytesIO()
    canvas.save(output, "JPEG", quality=92)
    return Response(output.getvalue(), media_type="image/jpeg", headers={"Cache-Control": "no-store", "X-Upscaled": "1" if used_ai else "0"})


@router.get("/card-printers")
def card_printers(refresh: bool = False) -> dict:
    if refresh:
        card_printing.invalidate_catalog()
    try:
        printers = card_printing.list_printers()
    except card_printing.CardPrintError as exc:
        raise HTTPException(503, str(exc)) from exc
    return {"printers": printers}


@router.post("/sessions/{session_id}/cards/print")
def print_cards(session_id: str, payload: CardPrintRequest, request: Request) -> dict:
    """Imprime en la impresora de tarjetas: una página por cara, en doble faz si corresponde."""
    require_same_origin(request)
    session = STORE.get(session_id)
    design = payload.design
    sides = requested_sides(design, payload.sides)
    records = [(session_record(session, item.photo_id), item) for item in payload.items]
    if len({item.photo_id for item in payload.items}) != len(payload.items):
        raise HTTPException(400, "El lote contiene credenciales repetidas")
    encode = payload.encode and design.magstripe.enabled
    encodings: dict[str, list[str]] = {}
    if encode:
        if not any(payload.mag_start):
            raise HTTPException(422, "Faltan los comandos de codificación de la impresora")
        for record, item in records:
            values = record_values(item.metadata, record.filename, item.sequence)
            tracks = magstripe_tracks(design, values)
            for track, data in tracks.items():
                error = card_magstripe.track_error(track, data)
                if error:
                    raise HTTPException(422, f"{record.filename}: {error}")
            strings = card_magstripe.encoding_strings(tracks, payload.mag_start, payload.mag_end)
            if not strings:
                raise HTTPException(422, f"{record.filename}: no hay datos para la banda magnética")
            encodings[item.photo_id] = strings
    landscape = design.width_mm >= design.height_mm
    title = f"Credenciales ({len(records)})" if len(records) > 1 else Path(records[0][0].filename).stem
    printed = []
    log_rows = []
    if not card_printing.PRINT_LOCK.acquire(timeout=5):
        raise HTTPException(409, "Ya hay una impresión de credenciales en curso")
    try:
        with card_printing.CardPrintJob(payload.printer, landscape=landscape, duplex=len(sides) == 2, title=title) as job:
            renderer = card_renderer(session, design, min(600, max(150, job.dpi)))
            for record, item in records:
                faces = render_card_faces(renderer, design, record, item, payload, sides)
                for _copy in range(payload.copies):
                    job.add_card(faces, rotate_back=payload.rotate_back, encode=encodings.get(item.photo_id))
                stamp = card_printing.timestamp()
                printed.append({"photo_id": item.photo_id, "printed_at": stamp, "encoded": item.photo_id in encodings})
                summary = " · ".join(value for value in list(item.metadata.values())[:4] if value)
                log_rows.append([stamp, payload.printer, record.filename, summary, "+".join(sides), str(payload.copies), "sí" if item.photo_id in encodings else ""])
            pages = job.pages
    except card_printing.CardPrintError as exc:
        raise HTTPException(502, str(exc)) from exc
    finally:
        card_printing.PRINT_LOCK.release()
    try:
        card_printing.append_print_log(log_rows)
    except OSError:
        pass
    return {"printer": payload.printer, "cards": len(printed) * payload.copies, "pages": pages, "printed": printed, "encoded": len(encodings)}


@router.get("/magstripe/profiles")
def magstripe_profiles() -> dict:
    """Comandos de codificación magnética de cada marca, para el diálogo de impresión."""
    return {"profiles": card_magstripe.public_profiles()}


@router.post("/sessions/{session_id}/project")
def save_project(session_id: str, payload: ProjectRequest, request: Request) -> StreamingResponse:
    require_same_origin(request)
    session = STORE.get(session_id)
    photo_states = payload.state.get("photos", {}) if isinstance(payload.state.get("photos", {}), dict) else {}
    manifest = []
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=4) as bundle:
        for record in session.photos.values():
            extension = record.source_path.suffix.casefold()
            member = f"photos/{record.id}{extension}"
            bundle.write(record.source_path, member)
            manifest.append({
                "archive_name": member,
                "filename": record.filename,
                "photo_id": record.id,
                "state": photo_states.get(record.id, {}),
            })
        assets = []
        for asset in session.assets.values():
            if not asset.path.exists():
                continue
            member = f"assets/{asset.id}{asset.path.suffix.casefold()}"
            bundle.write(asset.path, member)
            assets.append({"id": asset.id, "filename": asset.filename, "archive_name": member})
        project = {
            "version": 3,
            "saved_at": datetime.now().isoformat(timespec="seconds"),
            "state": payload.state,
            "manifest": manifest,
            "assets": assets,
        }
        encoded_project = json.dumps(project, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
        if len(encoded_project) > 5 * 1024 * 1024:
            raise HTTPException(413, "Los datos del proyecto son demasiado grandes")
        bundle.writestr("project.json", encoded_project)
    archive.seek(0)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return StreamingResponse(archive, media_type="application/zip", headers=attachment_headers(f"proyecto_fotos_{stamp}.recortador.zip"))


def restore_project_assets(session: PhotoSession, project: dict, bundle: zipfile.ZipFile, names: set[str]) -> None:
    entries = list(project.get("assets") or []) if project.get("version") == 3 else []
    if project.get("version") == 2:
        # Proyectos anteriores: logo y fondo fijos.
        for asset_id, key in (("logo", "logo"), ("card_background", "card_background")):
            member = project.get(key)
            if isinstance(member, str):
                entries.append({"id": asset_id, "filename": Path(member).name, "archive_name": member})
    for entry in entries[:MAX_ASSETS]:
        if not isinstance(entry, dict):
            continue
        asset_id = str(entry.get("id", ""))
        member = str(entry.get("archive_name", ""))
        extension = Path(member).suffix.casefold()
        if (
            not asset_id or len(asset_id) > 40 or not all(char.isalnum() or char == "_" for char in asset_id)
            or asset_id != asset_id.casefold() or member not in names or not member.startswith("assets/")
            or extension not in ASSET_EXTENSIONS or bundle.getinfo(member).file_size > MAX_ASSET_BYTES
        ):
            continue
        path = session.folder / "assets" / f"{asset_id}{extension}"
        path.write_bytes(bundle.read(member))
        session.assets[asset_id] = SessionAsset(asset_id, Path(str(entry.get("filename") or member)).name[:220], path)


@router.post("/projects", status_code=201)
async def open_project(
    request: Request,
    filename: str = Query(..., min_length=1, max_length=500),
) -> dict:
    require_same_origin(request)
    if not filename.casefold().endswith((".zip", ".recortador")):
        raise HTTPException(415, "El proyecto debe ser un archivo ZIP del recortador")
    content = await read_limited_body(request, MAX_PROJECT_BYTES, "El proyecto")
    session = STORE.create()
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as bundle:
            names = set(bundle.namelist())
            if "project.json" not in names:
                raise HTTPException(415, "El archivo no contiene un proyecto valido")
            if len(names) > MAX_PHOTOS + MAX_ASSETS + 10 or sum(info.file_size for info in bundle.infolist()) > MAX_SESSION_BYTES + MAX_ASSETS * MAX_ASSET_BYTES:
                raise HTTPException(413, "El proyecto supera los limites permitidos")
            project_info = bundle.getinfo("project.json")
            if project_info.file_size > 5 * 1024 * 1024:
                raise HTTPException(413, "Los datos del proyecto son demasiado grandes")
            project = json.loads(bundle.read("project.json").decode("utf-8"))
            if project.get("version") not in {2, 3} or not isinstance(project.get("manifest"), list):
                raise HTTPException(415, "Version de proyecto no admitida")
            restored = []
            for item in project["manifest"]:
                member = str(item.get("archive_name", ""))
                name = Path(str(item.get("filename", "foto.jpg")).replace("\\", "/")).name[:220]
                extension = Path(member).suffix.casefold()
                if member not in names or not member.startswith("photos/") or extension not in SUPPORTED_EXTENSIONS:
                    raise HTTPException(415, "El proyecto contiene una foto invalida")
                info = bundle.getinfo(member)
                if info.file_size > MAX_FILE_BYTES:
                    raise HTTPException(413, "Una foto del proyecto supera el limite permitido")
                identifier = uuid.uuid4().hex
                source_path = session.folder / "source" / f"{identifier}{extension}"
                source_path.write_bytes(bundle.read(member))
                record = register_source_image(session, source_path, name, photo_id=identifier)
                restored.append({**record.public(session.id), "state": item.get("state", {})})
            restore_project_assets(session, project, bundle, names)
            return {
                "session_id": session.id,
                "version": project.get("version"),
                "photos": restored,
                "state": project.get("state", {}),
                "assets": [asset.public(session.id) for asset in session.assets.values()],
                "logo_url": f"/api/photo-crop/sessions/{session.id}/logo" if session.logo_path else None,
                "card_background_url": f"/api/photo-crop/sessions/{session.id}/card-background" if session.card_background_path else None,
            }
    except HTTPException:
        STORE.delete(session.id)
        raise
    except (zipfile.BadZipFile, KeyError, json.JSONDecodeError, UnicodeDecodeError, OSError, ValueError) as exc:
        STORE.delete(session.id)
        raise HTTPException(415, "No se pudo abrir el proyecto") from exc


@router.delete("/sessions/{session_id}", status_code=204)
def delete_session(session_id: str, request: Request) -> Response:
    require_same_origin(request)
    STORE.delete(session_id)
    return Response(status_code=204)
