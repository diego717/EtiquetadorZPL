import asyncio
import io
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

project_root = Path(__file__).parent.parent
sys.path.insert(0, str(project_root / "src"))
sys.path.insert(0, str(project_root / "api"))

import sales_projection as projection_module
import sales_projection_endpoints as endpoints
from odoo_integration import OdooIntegration
from sales_projection import SalesProjectionMonitor, compute_rows, month_window, normalize_config, summarize

TODAY = date(2026, 10, 9)
WINDOW = month_window(TODAY)
LAST_12 = WINDOW["complete"][-12:]


def _product(product_id=1, on_hand=0.0, projected=None, outgoing=0.0, code="", category="CINTAS"):
    return {
        "id": product_id,
        "default_code": code or f"P{product_id}",
        "name": f"Producto {product_id}",
        "category": category,
        "qty_available": on_hand,
        "virtual_available": on_hand - outgoing if projected is None else projected,
        "incoming_qty": 0.0,
        "outgoing_qty": outgoing,
        "list_price": 10.0,
    }


def _flat_sales(qty, months=LAST_12, price=10.0):
    return {month: {"qty": float(qty), "amount": float(qty) * price} for month in months}


def _config(**overrides):
    base = {"use_seasonality": False, "lead_time_days": 5, "service_level": 95, "base_months": 6, "target_coverage_days": 30}
    base.update(overrides)
    return normalize_config(base)


def _row(product, sales=None, productions=None, rules=None, config=None):
    return compute_rows([product], {product["id"]: sales or {}}, productions or {}, rules, config or _config(), TODAY)[0]


class TestMonthWindow(unittest.TestCase):
    def test_forecast_starts_with_current_month(self):
        self.assertEqual(WINDOW["current"], "2026-10")
        self.assertEqual(WINDOW["forecast"], ["2026-10", "2026-11", "2026-12"])
        self.assertEqual(WINDOW["display"][0], "2025-10")
        self.assertEqual(WINDOW["display"][-1], "2026-09")
        self.assertEqual(WINDOW["since"], "2024-10-01")

    def test_year_rollover(self):
        window = month_window(date(2027, 1, 15))
        self.assertEqual(window["display"][-1], "2026-12")
        self.assertEqual(window["forecast"], ["2027-01", "2027-02", "2027-03"])


class TestComputeRows(unittest.TestCase):
    def test_flat_demand_gives_level_and_no_safety_stock(self):
        row = _row(_product(on_hand=1000), _flat_sales(304))
        self.assertEqual(row["forecast"][0]["qty"], 304.0)
        self.assertEqual(row["forecast"][0]["amount"], 3040.0)
        self.assertEqual(row["daily_demand"], 10.0)
        self.assertEqual(row["safety_stock"], 0.0)
        # 10/dia x 5 dias de lead time; maximo agrega 30 dias.
        self.assertEqual(row["min_suggested"], 50.0)
        self.assertEqual(row["max_suggested"], 350.0)
        self.assertEqual(row["status"], "ok")
        self.assertEqual(row["coverage_days"], 100.0)

    def test_weighted_average_favours_recent_months(self):
        sales = _flat_sales(100)
        sales[LAST_12[-1]] = {"qty": 400.0, "amount": 4000.0}
        row = _row(_product(on_hand=0), sales)
        # Pesos 1..6: (100*15 + 400*6) / 21
        self.assertAlmostEqual(row["monthly_level"], 185.71, places=2)

    def test_variable_demand_adds_safety_stock(self):
        sales = {month: {"qty": 0.0 if i % 2 else 608.0, "amount": 0.0} for i, month in enumerate(LAST_12)}
        row = _row(_product(on_hand=0), sales)
        self.assertGreater(row["safety_stock"], 0)
        self.assertGreater(row["min_suggested"], row["daily_demand"] * 5)

    def test_low_stock_suggests_up_to_maximum(self):
        row = _row(_product(on_hand=30), _flat_sales(304))
        self.assertEqual(row["status"], "fabricar")
        self.assertEqual(row["suggested_qty"], 320.0)
        self.assertEqual(row["suggested_amount"], 3200.0)
        self.assertEqual(row["stockout_date"], "2026-10-12")

    def test_watch_status_inside_one_extra_lead_time(self):
        row = _row(_product(on_hand=80), _flat_sales(304))
        self.assertEqual(row["status"], "atento")
        self.assertEqual(row["suggested_qty"], 0.0)

    def test_negative_stock_is_treated_as_zero_and_flagged(self):
        row = _row(_product(on_hand=-500, outgoing=20), _flat_sales(304))
        self.assertTrue(row["negative_stock"])
        self.assertEqual(row["virtual_available"], -20.0)
        self.assertEqual(row["status"], "quiebre")
        self.assertEqual(row["suggested_qty"], 370.0)

    def test_draft_production_counts_in_projection(self):
        productions = {1: {"qty": 300.0, "draft_qty": 300.0, "orders": ["MO/1"], "next_date": "2026-10-10"}}
        row = _row(_product(on_hand=30), _flat_sales(304), productions)
        self.assertEqual(row["in_production"], 300.0)
        self.assertEqual(row["virtual_available"], 330.0)
        self.assertEqual(row["status"], "ok")

    def test_sporadic_product_is_make_to_order(self):
        sales = {LAST_12[-1]: {"qty": 5000.0, "amount": 70000.0}}
        row = _row(_product(on_hand=0), sales)
        self.assertEqual(row["mode"], "pedido")
        self.assertEqual(row["mode_source"], "auto")
        self.assertEqual(row["status"], "a_pedido")
        self.assertEqual(row["min_suggested"], 0.0)
        self.assertEqual(row["suggested_qty"], 0.0)

    def test_make_to_order_with_uncovered_orders_suggests_only_the_gap(self):
        sales = {LAST_12[-1]: {"qty": 5000.0, "amount": 70000.0}}
        row = _row(_product(on_hand=100, outgoing=600), sales)
        self.assertEqual(row["status"], "pedido_pendiente")
        self.assertEqual(row["suggested_qty"], 500.0)

    def test_manual_overrides(self):
        product = _product(code="NEO106COLL", category="NEOCARD / PINES / PERSONALIZADOS")
        by_category = _row(product, _flat_sales(304), config=_config(make_to_order_categories=["personalizados"]))
        self.assertEqual((by_category["mode"], by_category["mode_source"]), ("pedido", "categoria"))
        forced_stock = _row(
            product,
            _flat_sales(304),
            config=_config(make_to_order_categories=["PERSONALIZADOS"], make_to_stock_codes=["neo106coll"]),
        )
        self.assertEqual((forced_stock["mode"], forced_stock["mode_source"]), ("stock", "producto"))

    def test_seasonality_needs_regular_sales_and_is_clamped(self):
        sales = _flat_sales(100, WINDOW["complete"])
        sales["2025-10"] = {"qty": 1000.0, "amount": 10000.0}
        row = _row(_product(on_hand=0), sales, config=_config(use_seasonality=True))
        # oct-25 fue 1000 contra un promedio anual de 175: 5,7 se recorta a 2.
        self.assertEqual(row["forecast"][0]["seasonal_factor"], 2.0)
        self.assertEqual(row["forecast"][0]["qty"], 200.0)
        # nov-25 fue 100 contra el mismo promedio de 175.
        self.assertEqual(row["forecast"][1]["seasonal_factor"], 0.57)

        sparse = {month: values for month, values in sales.items() if month >= "2025-10"}
        sparse_row = _row(_product(on_hand=0), sparse, config=_config(use_seasonality=True))
        self.assertIsNone(sparse_row["forecast"][0]["seasonal_factor"])

    def test_without_sales_uses_list_price_and_is_idle(self):
        row = _row(_product(on_hand=10))
        self.assertEqual(row["status"], "sin_venta")
        self.assertEqual(row["avg_price"], 10.0)
        self.assertEqual(row["abc"], "-")

    def test_abc_ranking_and_summary(self):
        products = [_product(i, on_hand=10000) for i in (1, 2, 3)]
        sales = {1: _flat_sales(100, price=80), 2: _flat_sales(100, price=15), 3: _flat_sales(100, price=5)}
        rows = compute_rows(products, sales, {}, None, _config(), TODAY)
        self.assertEqual({row["product_id"]: row["abc"] for row in rows}, {1: "A", 2: "B", 3: "C"})
        summary = summarize(rows)
        self.assertEqual(summary["total"], 3)
        self.assertEqual(summary["amount_12m"], 120000.0)

    def test_odoo_rule_minimum_is_reported(self):
        row = _row(_product(on_hand=1000), _flat_sales(304), rules={1: {"min_qty": 25.0, "max_qty": 0.0, "rules": 1}})
        self.assertEqual(row["odoo_min_qty"], 25.0)
        self.assertIsNone(_row(_product(on_hand=1000), _flat_sales(304))["odoo_min_qty"])

    def test_config_is_normalized(self):
        config = normalize_config(
            {"service_level": 97, "lead_time_days": 0, "make_to_order_codes": "A1, b2,\nA1", "unknown": 1}
        )
        self.assertEqual(config["service_level"], 98)
        self.assertEqual(config["lead_time_days"], 1)
        self.assertEqual(config["make_to_order_codes"], ["A1", "b2"])
        self.assertNotIn("unknown", config)


class _FakeOdoo:
    def __init__(self):
        self.fail = False
        self.sales_since = None

    def is_configured(self):
        return True

    def humanize_exception(self, exc):
        return str(exc)

    def get_manufactured_products(self):
        if self.fail:
            raise ValueError("Odoo caido")
        return [_product(1, on_hand=30), _product(2, on_hand=1000), _product(3, on_hand=0)]

    def get_monthly_sales(self, product_ids, since_date):
        self.sales_since = since_date
        return {1: _flat_sales(304), 2: _flat_sales(304)}

    def get_open_productions(self):
        return {}

    def get_reorder_rules(self):
        return None


class TestMonitorAndEndpoints(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        tmp_path = Path(self.tmp.name)
        self.patcher = patch.object(projection_module, "_resolve_state_path", lambda name: tmp_path / name)
        self.patcher.start()
        self.odoo = _FakeOdoo()
        self.monitor = SalesProjectionMonitor(odoo=self.odoo)

    def tearDown(self):
        self.patcher.stop()
        self.tmp.cleanup()

    def test_refresh_saves_snapshot(self):
        snapshot = self.monitor.refresh(today=TODAY)
        self.assertEqual(self.odoo.sales_since, "2024-10-01")
        self.assertEqual(snapshot["forecast_months"], ["2026-10", "2026-11", "2026-12"])
        self.assertEqual(snapshot["summary"]["fabricar"], 1)
        self.assertEqual(self.monitor.load_snapshot()["generated_at"], snapshot["generated_at"])

    def test_failed_refresh_keeps_previous_rows(self):
        self.monitor.refresh(today=TODAY)
        self.odoo.fail = True
        with self.assertRaises(ValueError):
            self.monitor.refresh(today=TODAY)
        snapshot = self.monitor.load_snapshot()
        self.assertEqual(len(snapshot["rows"]), 3)
        self.assertIn("Odoo caido", snapshot["last_error"])

    def test_report_hides_idle_products_unless_asked(self):
        self.monitor.refresh(today=TODAY)
        with patch.object(endpoints, "sales_projection_monitor", self.monitor):
            default = asyncio.run(endpoints.get_report(limit=5000))
            everything = asyncio.run(endpoints.get_report(include_idle=True, limit=5000))
            alerts = asyncio.run(endpoints.get_report(only_alerts=True, limit=5000))
        self.assertEqual(sorted(row["product_id"] for row in default["items"]), [1, 2])
        self.assertEqual(everything["total"], 3)
        self.assertEqual([row["product_id"] for row in alerts["items"]], [1])

    def test_export_builds_workbook_with_pivot_sheet(self):
        from openpyxl import load_workbook

        snapshot = self.monitor.refresh(today=TODAY)
        content = endpoints.build_workbook(snapshot, snapshot["rows"])
        wb = load_workbook(io.BytesIO(content))
        self.assertEqual(wb.sheetnames, ["Proyeccion", "Datos", "Parametros"])
        self.assertEqual(wb["Proyeccion"].max_row, 4)
        # Por producto: 12 meses reales + mes en curso + 3 proyectados.
        self.assertEqual(wb["Datos"].max_row, 1 + 3 * 16)
        self.assertIn("Proy. oct-26 (u)", [cell.value for cell in wb["Proyeccion"][1]])


class TestMonthlySalesQuery(unittest.TestCase):
    def test_groups_by_local_month_and_converts_currency(self):
        lines = [
            # 2026-10-01 02:00 UTC = 30/09 23:00 en Uruguay.
            {"product_id": [7, "X"], "product_uom_qty": 10, "price_subtotal": 1000.0, "order_id": [1, "S1"]},
            {"product_id": [7, "X"], "product_uom_qty": 5, "price_subtotal": 50.0, "order_id": [2, "S2"]},
        ]
        orders = [
            {"id": 1, "date_order": "2026-10-01 02:00:00", "currency_rate": 1.0},
            {"id": 2, "date_order": "2026-10-05 15:00:00", "currency_rate": 0.025},
        ]
        domains = []

        def fake_call(model, method, args, kwargs=None):
            if model == "sale.order.line":
                domains.append(args[0])
                return lines if not (kwargs or {}).get("offset") else []
            if model == "sale.order":
                return orders
            raise AssertionError(model)

        odoo = OdooIntegration.__new__(OdooIntegration)
        with patch.object(OdooIntegration, "_readonly_session", lambda self, auth=None: (1, fake_call)):
            result = odoo.get_monthly_sales([7], "2024-10-01")
        self.assertEqual(result[7]["2026-09"], {"qty": 10.0, "amount": 1000.0})
        self.assertEqual(result[7]["2026-10"], {"qty": 5.0, "amount": 2000.0})
        self.assertIn(("order_id.date_order", ">=", "2024-10-01 03:00:00"), domains[0])


if __name__ == "__main__":
    unittest.main()
