"""
Genera un Excel con informes contables de Odoo para un periodo.

Hojas: Resumen, Ventas por mes, Compras por mes, Estado de resultados,
Sumas y saldos, Libro de ventas, Libro de compras, IVA por tasa,
Cuentas por cobrar, Cuentas por pagar y Cobros y pagos.

Solo considera asientos publicados (state = posted); los borradores quedan
fuera. Los importes van en la moneda de la compania (UYU).

Usa la configuracion ya guardada por la app (odoo_config.json).

Ejemplos:
    python scripts/informes_contables_odoo.py
    python scripts/informes_contables_odoo.py --desde 2026-01-01 --hasta 2026-06-30
    python scripts/informes_contables_odoo.py --salida "C:/Temp/contabilidad.xlsx"
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT / "config"))

from odoo_integration import odoo_integration  # noqa: E402

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - dependencia opcional
    print("Falta openpyxl. Instalalo con: pip install openpyxl")
    raise SystemExit(1)

MONEY = "#,##0.00"
INT = "#,##0"
BATCH = 500

INCOME_TYPES = ["income", "income_other"]
EXPENSE_TYPES = ["expense", "expense_direct_cost", "expense_depreciation"]

ACCOUNT_TYPE_LABELS = {
    "asset_receivable": "Deudores comerciales",
    "asset_cash": "Caja y bancos",
    "asset_current": "Activo corriente",
    "asset_fixed": "Activo fijo",
    "liability_payable": "Acreedores comerciales",
    "liability_current": "Pasivo corriente",
    "equity": "Patrimonio",
    "equity_unaffected": "Resultados no asignados",
    "income": "Ingresos operativos",
    "income_other": "Otros ingresos",
    "expense": "Gastos",
    "expense_direct_cost": "Costo de ventas",
    "expense_depreciation": "Amortizaciones",
}


class Odoo:
    """Envoltorio minimo de XML-RPC sobre la configuracion de la app."""

    def __init__(self) -> None:
        self.uid = odoo_integration._authenticate()
        self.models = odoo_integration._xmlrpc_models()
        self.db = odoo_integration.config.get("database", "")
        self.password = odoo_integration.config.get("password", "")

    def call(self, model: str, method: str, args: List[Any], kwargs: Optional[Dict[str, Any]] = None) -> Any:
        return self.models.execute_kw(self.db, self.uid, self.password, model, method, args, kwargs or {})

    def read_group(self, model: str, domain: List[Any], fields: List[str], groupby: List[str]) -> List[Dict[str, Any]]:
        return self.call(model, "read_group", [domain, fields, groupby], {"lazy": False})

    def search_read(self, model: str, domain: List[Any], fields: List[str], order: str = "") -> List[Dict[str, Any]]:
        total = self.call(model, "search_count", [domain])
        registros: List[Dict[str, Any]] = []
        offset = 0
        while offset < total:
            opciones: Dict[str, Any] = {"fields": fields, "offset": offset, "limit": BATCH}
            if order:
                opciones["order"] = order
            lote = self.call(model, "search_read", [domain], opciones)
            if not lote:
                break
            registros.extend(lote)
            offset += len(lote)
        return registros

    def read(self, model: str, ids: List[int], fields: List[str]) -> List[Dict[str, Any]]:
        salida: List[Dict[str, Any]] = []
        ids = list(ids)
        for inicio in range(0, len(ids), 1000):
            salida.extend(self.call(model, "read", [ids[inicio:inicio + 1000]], {"fields": fields}))
        return salida


def _m2o_name(value: Any) -> str:
    if isinstance(value, (list, tuple)) and len(value) >= 2:
        return str(value[1] or "").strip()
    return ""


def _m2o_id(value: Any) -> Optional[int]:
    if isinstance(value, (list, tuple)) and value:
        return int(value[0])
    return None


def _text(value: Any) -> str:
    if value is False or value is None:
        return ""
    return str(value).strip()


# ---------------------------------------------------------------- extraccion

def cargar_datos(odoo: Odoo, desde: str, hasta: str) -> Dict[str, Any]:
    periodo_mov = [("date", ">=", desde), ("date", "<=", hasta), ("parent_state", "=", "posted")]
    periodo_fact = [("date", ">=", desde), ("date", "<=", hasta), ("state", "=", "posted")]
    datos: Dict[str, Any] = {}

    print("  plan de cuentas...")
    cuentas = odoo.search_read("account.account", [], ["code", "name", "account_type"], order="code asc")
    datos["cuentas"] = {c["id"]: c for c in cuentas}

    print("  sumas y saldos...")
    datos["saldos"] = odoo.read_group("account.move.line", periodo_mov, ["debit", "credit", "balance"], ["account_id"])

    print("  ventas y compras por mes...")
    datos["ventas_mes"] = odoo.read_group(
        "account.move",
        periodo_fact + [("move_type", "in", ["out_invoice", "out_refund"])],
        ["amount_untaxed_signed", "amount_tax_signed", "amount_total_signed"],
        ["date:month"],
    )
    datos["compras_mes"] = odoo.read_group(
        "account.move",
        periodo_fact + [("move_type", "in", ["in_invoice", "in_refund"])],
        ["amount_untaxed_signed", "amount_tax_signed", "amount_total_signed"],
        ["date:month"],
    )

    print("  impuestos por tasa...")
    datos["impuestos"] = odoo.read_group(
        "account.move.line",
        periodo_mov + [("tax_line_id", "!=", False)],
        ["balance"],
        ["tax_line_id", "journal_id"],
    )

    campos_factura = [
        "name", "date", "invoice_date", "invoice_date_due", "partner_id", "move_type",
        "amount_untaxed_signed", "amount_tax_signed", "amount_total_signed",
        "amount_total", "amount_residual_signed", "currency_id", "payment_state",
        "journal_id", "invoice_origin", "ref",
    ]
    print("  libro de ventas...")
    datos["ventas"] = odoo.search_read(
        "account.move", periodo_fact + [("move_type", "in", ["out_invoice", "out_refund"])],
        campos_factura, order="date asc, name asc",
    )
    print("  libro de compras...")
    datos["compras"] = odoo.search_read(
        "account.move", periodo_fact + [("move_type", "in", ["in_invoice", "in_refund"])],
        campos_factura, order="date asc, name asc",
    )

    # Odoo firma las compras desde la optica de la empresa (una factura de
    # proveedor es negativa). Las invertimos para que el libro de compras se
    # lea con los mismos signos que el de ventas.
    for registros in (datos["compras"], datos["compras_mes"]):
        for registro in registros:
            for campo in ("amount_untaxed_signed", "amount_tax_signed",
                          "amount_total_signed", "amount_residual_signed"):
                if campo in registro and registro[campo] is not None:
                    registro[campo] = -registro[campo]

    print("  desglose de IVA por comprobante...")
    lineas_iva = odoo.search_read(
        "account.move.line", periodo_mov + [("tax_line_id", "!=", False)],
        ["move_id", "tax_line_id", "balance"],
    )
    por_move: Dict[int, Dict[str, float]] = defaultdict(lambda: defaultdict(float))
    for linea in lineas_iva:
        move = _m2o_id(linea.get("move_id"))
        if move is None:
            continue
        por_move[move][_m2o_name(linea.get("tax_line_id"))] += float(linea.get("balance") or 0.0)
    datos["iva_por_move"] = por_move

    print("  saldos abiertos de clientes y proveedores...")
    abierto = [("date", "<=", hasta), ("parent_state", "=", "posted"),
               ("full_reconcile_id", "=", False), ("amount_residual", "!=", 0)]
    campos_abierto = ["partner_id", "move_name", "date", "date_maturity", "amount_residual", "account_id"]
    datos["por_cobrar"] = odoo.search_read(
        "account.move.line", abierto + [("account_id.account_type", "=", "asset_receivable")], campos_abierto)
    datos["por_pagar"] = odoo.search_read(
        "account.move.line", abierto + [("account_id.account_type", "=", "liability_payable")], campos_abierto)

    print("  cobros y pagos...")
    datos["pagos"] = odoo.read_group(
        "account.payment",
        [("date", ">=", desde), ("date", "<=", hasta), ("state", "not in", ["draft", "cancel"])],
        ["amount_company_currency_signed"],
        ["journal_id", "payment_type"],
    )

    print("  identificacion de clientes y proveedores...")
    ids_socios = set()
    for factura in datos["ventas"] + datos["compras"]:
        pid = _m2o_id(factura.get("partner_id"))
        if pid:
            ids_socios.add(pid)
    for linea in datos["por_cobrar"] + datos["por_pagar"]:
        pid = _m2o_id(linea.get("partner_id"))
        if pid:
            ids_socios.add(pid)
    socios = odoo.read("res.partner", sorted(ids_socios),
                       ["name", "vat", "l10n_latam_identification_type_id"])
    datos["socios"] = {s["id"]: s for s in socios}
    return datos


def _doc_socio(datos: Dict[str, Any], partner: Any) -> Tuple[str, str]:
    socio = datos["socios"].get(_m2o_id(partner) or -1)
    if not socio:
        return "", ""
    return _m2o_name(socio.get("l10n_latam_identification_type_id")), _text(socio.get("vat"))


def _iva_desglose(datos: Dict[str, Any], move_id: int, signo: int) -> Dict[str, float]:
    return {nombre: importe * signo for nombre, importe in datos["iva_por_move"].get(move_id, {}).items()}


# -------------------------------------------------------------------- estilos

TITULO_FILL = PatternFill("solid", fgColor="1F4E78")
TITULO_FONT = Font(color="FFFFFF", bold=True)
TOTAL_FONT = Font(bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
BORDE_SUP = Border(top=Side(style="thin", color="808080"))

Columna = Tuple[str, int, str]


def nueva_hoja(workbook: Workbook, titulo: str, columnas: List[Columna]):
    hoja = workbook.create_sheet()
    hoja.title = titulo
    hoja.append([nombre for nombre, _, _ in columnas])
    for celda in hoja[1]:
        celda.fill = TITULO_FILL
        celda.font = TITULO_FONT
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    for indice, (_, ancho, _) in enumerate(columnas, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.freeze_panes = "A2"
    return hoja


def volcar(hoja, columnas: List[Columna], filas: Iterable[List[Any]]) -> None:
    for fila in filas:
        hoja.append(fila)
        for indice, (_, _, formato) in enumerate(columnas, start=1):
            if formato:
                hoja.cell(row=hoja.max_row, column=indice).number_format = formato


def fila_total(hoja, columnas: List[Columna], etiqueta: str, valores: Dict[int, float]) -> None:
    fila: List[Any] = [""] * len(columnas)
    fila[0] = etiqueta
    for indice, valor in valores.items():
        fila[indice] = valor
    hoja.append(fila)
    for indice, (_, _, formato) in enumerate(columnas, start=1):
        celda = hoja.cell(row=hoja.max_row, column=indice)
        celda.font = TOTAL_FONT
        celda.fill = TOTAL_FILL
        celda.border = BORDE_SUP
        if formato:
            celda.number_format = formato


# --------------------------------------------------------------------- hojas

def hoja_resumen(workbook: Workbook, datos: Dict[str, Any], desde: str, hasta: str) -> None:
    ventas_neto = sum(f["amount_untaxed_signed"] for f in datos["ventas"])
    ventas_iva = sum(f["amount_tax_signed"] for f in datos["ventas"])
    ventas_total = sum(f["amount_total_signed"] for f in datos["ventas"])
    compras_neto = sum(f["amount_untaxed_signed"] for f in datos["compras"])
    compras_iva = sum(f["amount_tax_signed"] for f in datos["compras"])
    compras_total = sum(f["amount_total_signed"] for f in datos["compras"])

    ingresos = gastos = 0.0
    for grupo in datos["saldos"]:
        cuenta = datos["cuentas"].get(_m2o_id(grupo.get("account_id")) or -1)
        if not cuenta:
            continue
        if cuenta["account_type"] in INCOME_TYPES:
            ingresos += -grupo["balance"]
        elif cuenta["account_type"] in EXPENSE_TYPES:
            gastos += grupo["balance"]

    por_cobrar = sum(l["amount_residual"] for l in datos["por_cobrar"])
    por_pagar = -sum(l["amount_residual"] for l in datos["por_pagar"])

    hoja = workbook.active
    hoja.title = "Resumen"
    hoja.column_dimensions["A"].width = 42
    hoja.column_dimensions["B"].width = 22

    hoja.append(["Informes contables - Odoo"])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append(["Periodo", f"{desde} a {hasta}"])
    hoja.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M")])
    hoja.append(["Alcance", "Solo asientos publicados (posted)"])
    hoja.append(["Moneda", "UYU (moneda de la compania)"])
    hoja.append([])

    def bloque(titulo: str, filas: List[Tuple[str, float, str]]) -> None:
        hoja.append([titulo])
        celda = hoja.cell(row=hoja.max_row, column=1)
        celda.font = TOTAL_FONT
        celda.fill = TOTAL_FILL
        for etiqueta, valor, formato in filas:
            hoja.append([etiqueta, valor])
            hoja.cell(row=hoja.max_row, column=2).number_format = formato
        hoja.append([])

    bloque("Ventas del periodo", [
        ("Comprobantes emitidos", len(datos["ventas"]), INT),
        ("Neto (gravado + exento)", ventas_neto, MONEY),
        ("IVA ventas", ventas_iva, MONEY),
        ("Total facturado", ventas_total, MONEY),
    ])
    bloque("Compras del periodo", [
        ("Comprobantes recibidos", len(datos["compras"]), INT),
        ("Neto", compras_neto, MONEY),
        ("IVA compras", compras_iva, MONEY),
        ("Total", compras_total, MONEY),
    ])
    bloque("IVA", [
        ("IVA ventas (debito)", ventas_iva, MONEY),
        ("IVA compras (credito)", compras_iva, MONEY),
        ("Diferencia (debito - credito)", ventas_iva - compras_iva, MONEY),
    ])
    bloque("Resultado del periodo", [
        ("Ingresos", ingresos, MONEY),
        ("Gastos y costos", gastos, MONEY),
        ("Resultado", ingresos - gastos, MONEY),
    ])
    bloque(f"Saldos abiertos al {hasta}", [
        ("A cobrar de clientes", por_cobrar, MONEY),
        ("A pagar a proveedores", por_pagar, MONEY),
    ])
    hoja.append(["Un saldo a pagar negativo indica pagos a proveedores todavia sin"])
    hoja.append(["conciliar contra su factura; se ven en la hoja Cuentas por pagar."])
    hoja.append([])
    hoja.append(["Nota: la diferencia de IVA es informativa y no reemplaza la liquidacion"])
    hoja.append(["de DGI: no contempla retenciones, anticipos ni ajustes del cierre."])


def hoja_por_mes(workbook: Workbook, datos: Dict[str, Any], clave: str, titulo: str) -> None:
    columnas: List[Columna] = [("Mes", 18, ""), ("Comprobantes", 14, INT), ("Neto", 16, MONEY),
                               ("IVA", 16, MONEY), ("Total", 18, MONEY)]
    hoja = nueva_hoja(workbook, titulo, columnas)
    filas = [[g["date:month"], g["__count"], g["amount_untaxed_signed"],
              g["amount_tax_signed"], g["amount_total_signed"]] for g in datos[clave]]
    volcar(hoja, columnas, filas)
    if filas:
        fila_total(hoja, columnas, "TOTAL", {indice: sum(f[indice] for f in filas) for indice in range(1, 5)})


def hoja_resultados(workbook: Workbook, datos: Dict[str, Any]) -> None:
    columnas: List[Columna] = [("Codigo", 12, ""), ("Cuenta", 46, ""), ("Categoria", 24, ""), ("Importe", 18, MONEY)]
    hoja = nueva_hoja(workbook, "Estado de resultados", columnas)

    por_tipo: Dict[str, List[List[Any]]] = defaultdict(list)
    for grupo in datos["saldos"]:
        cuenta = datos["cuentas"].get(_m2o_id(grupo.get("account_id")) or -1)
        if not cuenta or cuenta["account_type"] not in INCOME_TYPES + EXPENSE_TYPES:
            continue
        signo = -1 if cuenta["account_type"] in INCOME_TYPES else 1
        por_tipo[cuenta["account_type"]].append([
            cuenta["code"], cuenta["name"],
            ACCOUNT_TYPE_LABELS.get(cuenta["account_type"], cuenta["account_type"]),
            grupo["balance"] * signo,
        ])

    ingresos = gastos = 0.0
    for tipo in INCOME_TYPES + EXPENSE_TYPES:
        filas = sorted(por_tipo.get(tipo, []), key=lambda f: f[0] or "")
        if not filas:
            continue
        volcar(hoja, columnas, filas)
        subtotal = sum(f[3] for f in filas)
        fila_total(hoja, columnas, f"Subtotal {ACCOUNT_TYPE_LABELS.get(tipo, tipo)}", {3: subtotal})
        if tipo in INCOME_TYPES:
            ingresos += subtotal
        else:
            gastos += subtotal

    hoja.append([])
    fila_total(hoja, columnas, "TOTAL INGRESOS", {3: ingresos})
    fila_total(hoja, columnas, "TOTAL GASTOS Y COSTOS", {3: gastos})
    fila_total(hoja, columnas, "RESULTADO DEL PERIODO", {3: ingresos - gastos})


def hoja_sumas_saldos(workbook: Workbook, datos: Dict[str, Any]) -> None:
    columnas: List[Columna] = [("Codigo", 12, ""), ("Cuenta", 46, ""), ("Tipo", 26, ""),
                               ("Debe", 16, MONEY), ("Haber", 16, MONEY), ("Saldo", 16, MONEY)]
    hoja = nueva_hoja(workbook, "Sumas y saldos", columnas)
    filas = []
    for grupo in datos["saldos"]:
        cuenta = datos["cuentas"].get(_m2o_id(grupo.get("account_id")) or -1)
        if not cuenta:
            continue
        filas.append([cuenta["code"], cuenta["name"],
                      ACCOUNT_TYPE_LABELS.get(cuenta["account_type"], cuenta["account_type"]),
                      grupo["debit"], grupo["credit"], grupo["balance"]])
    filas.sort(key=lambda f: f[0] or "")
    volcar(hoja, columnas, filas)
    if filas:
        fila_total(hoja, columnas, "TOTAL", {indice: sum(f[indice] for f in filas) for indice in range(3, 6)})
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{hoja.max_row}"


def hoja_libro(workbook: Workbook, datos: Dict[str, Any], clave: str, titulo: str, etiqueta_socio: str) -> None:
    tasas = sorted({nombre for desglose in datos["iva_por_move"].values() for nombre in desglose})
    columnas: List[Columna] = [("Fecha", 12, ""), ("Comprobante", 22, ""), ("Tipo", 14, ""),
                               (etiqueta_socio, 36, ""), ("Tipo doc", 12, ""), ("RUT / Documento", 18, ""),
                               ("Neto", 15, MONEY)]
    columnas += [(f"IVA {tasa}", 13, MONEY) for tasa in tasas]
    columnas += [("IVA total", 15, MONEY), ("Total", 16, MONEY), ("Moneda", 9, ""),
                 ("Total moneda orig.", 16, MONEY), ("Estado de pago", 15, ""), ("Saldo pendiente", 15, MONEY)]
    hoja = nueva_hoja(workbook, titulo, columnas)

    etiquetas_tipo = {"out_invoice": "Factura", "out_refund": "Nota credito",
                      "in_invoice": "Factura", "in_refund": "Nota credito"}
    filas = []
    for factura in datos[clave]:
        signo = -1 if factura["move_type"] in ("out_invoice", "out_refund") else 1
        desglose = _iva_desglose(datos, factura["id"], signo)
        tipo_doc, documento = _doc_socio(datos, factura.get("partner_id"))
        fila: List[Any] = [
            _text(factura.get("invoice_date")) or _text(factura.get("date")),
            _text(factura.get("name")),
            etiquetas_tipo.get(factura["move_type"], factura["move_type"]),
            _m2o_name(factura.get("partner_id")),
            tipo_doc,
            documento,
            factura["amount_untaxed_signed"],
        ]
        fila += [desglose.get(tasa, 0.0) for tasa in tasas]
        fila += [
            factura["amount_tax_signed"],
            factura["amount_total_signed"],
            _m2o_name(factura.get("currency_id")),
            factura.get("amount_total") or 0.0,
            _text(factura.get("payment_state")),
            factura.get("amount_residual_signed") or 0.0,
        ]
        filas.append(fila)

    volcar(hoja, columnas, filas)
    if filas:
        # "Total moneda orig." queda fuera del total: mezcla UYU con USD.
        numericas = [indice for indice, (nombre, _, formato) in enumerate(columnas)
                     if formato == MONEY and nombre != "Total moneda orig."]
        fila_total(hoja, columnas, "TOTAL", {indice: sum(f[indice] for f in filas) for indice in numericas})
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{hoja.max_row}"


def hoja_iva(workbook: Workbook, datos: Dict[str, Any]) -> None:
    columnas: List[Columna] = [("Impuesto", 30, ""), ("Diario", 34, ""), ("Importe", 18, MONEY)]
    hoja = nueva_hoja(workbook, "IVA por tasa", columnas)
    filas = [[_m2o_name(g.get("tax_line_id")), _m2o_name(g.get("journal_id")), -g["balance"]]
             for g in datos["impuestos"]]
    filas.sort(key=lambda f: (f[0], f[1]))
    volcar(hoja, columnas, filas)
    if filas:
        fila_total(hoja, columnas, "TOTAL (positivo = debito fiscal)", {2: sum(f[2] for f in filas)})


def hoja_aging(workbook: Workbook, datos: Dict[str, Any], clave: str, titulo: str, corte: str, signo: int) -> None:
    columnas: List[Columna] = [("Socio", 40, ""), ("Tipo doc", 12, ""), ("RUT / Documento", 18, ""),
                               ("Al dia", 15, MONEY), ("1-30 dias", 15, MONEY), ("31-60 dias", 15, MONEY),
                               ("61-90 dias", 15, MONEY), ("+90 dias", 15, MONEY), ("Total", 16, MONEY)]
    hoja = nueva_hoja(workbook, titulo, columnas)
    corte_fecha = date.fromisoformat(corte)

    acumulado: Dict[int, List[float]] = defaultdict(lambda: [0.0] * 5)
    socios: Dict[int, Any] = {}
    for linea in datos[clave]:
        pid = _m2o_id(linea.get("partner_id")) or 0
        socios[pid] = linea.get("partner_id")
        importe = float(linea.get("amount_residual") or 0.0) * signo
        vencimiento = _text(linea.get("date_maturity")) or _text(linea.get("date"))
        try:
            dias = (corte_fecha - date.fromisoformat(vencimiento)).days
        except ValueError:
            dias = 0
        if dias <= 0:
            tramo = 0
        elif dias <= 30:
            tramo = 1
        elif dias <= 60:
            tramo = 2
        elif dias <= 90:
            tramo = 3
        else:
            tramo = 4
        acumulado[pid][tramo] += importe

    filas = []
    for pid, tramos in acumulado.items():
        tipo_doc, documento = _doc_socio(datos, socios.get(pid))
        filas.append([_m2o_name(socios.get(pid)) or "(sin socio)", tipo_doc, documento] + tramos + [sum(tramos)])
    filas.sort(key=lambda f: -f[-1])
    volcar(hoja, columnas, filas)
    if filas:
        fila_total(hoja, columnas, "TOTAL", {indice: sum(f[indice] for f in filas) for indice in range(3, 9)})
    hoja.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{hoja.max_row}"


def hoja_pagos(workbook: Workbook, datos: Dict[str, Any]) -> None:
    columnas: List[Columna] = [("Diario / medio", 40, ""), ("Sentido", 16, ""),
                               ("Comprobantes", 14, INT), ("Importe", 18, MONEY)]
    hoja = nueva_hoja(workbook, "Cobros y pagos", columnas)
    sentidos = {"inbound": "Cobros", "outbound": "Pagos"}
    filas = [[_m2o_name(g.get("journal_id")), sentidos.get(g.get("payment_type"), _text(g.get("payment_type"))),
              g["__count"], g["amount_company_currency_signed"]] for g in datos["pagos"]]
    filas.sort(key=lambda f: (f[1], -abs(f[3])))
    volcar(hoja, columnas, filas)
    if filas:
        fila_total(hoja, columnas, "TOTAL", {2: sum(f[2] for f in filas), 3: sum(f[3] for f in filas)})


# ----------------------------------------------------------------------- main

def main() -> int:
    hoy = date.today()
    parser = argparse.ArgumentParser(description="Genera informes contables de Odoo en Excel")
    parser.add_argument("--desde", default=f"{hoy.year}-01-01", help="Fecha inicial AAAA-MM-DD")
    parser.add_argument("--hasta", default=hoy.isoformat(), help="Fecha final AAAA-MM-DD")
    parser.add_argument("--salida", default="", help="Ruta del .xlsx a generar")
    args = parser.parse_args()

    for etiqueta, valor in (("--desde", args.desde), ("--hasta", args.hasta)):
        try:
            date.fromisoformat(valor)
        except ValueError:
            print(f"Fecha invalida en {etiqueta}: {valor} (formato AAAA-MM-DD)")
            return 1
    if args.desde > args.hasta:
        print("Rango invalido: --desde es posterior a --hasta")
        return 1

    if not odoo_integration.is_configured():
        print("Odoo no esta configurado. Completa URL, base de datos, usuario y clave en la app.")
        return 1

    if args.salida:
        destino = Path(args.salida).expanduser()
    else:
        destino = Path.home() / "Desktop" / f"informes_contables_{args.desde}_a_{args.hasta}.xlsx"

    print(f"Extrayendo datos de Odoo ({args.desde} a {args.hasta})...")
    try:
        odoo = Odoo()
        datos = cargar_datos(odoo, args.desde, args.hasta)
    except Exception as exc:
        print(f"Error consultando Odoo: {odoo_integration.humanize_exception(exc)}")
        return 1

    print("Armando el Excel...")
    workbook = Workbook()
    hoja_resumen(workbook, datos, args.desde, args.hasta)
    hoja_por_mes(workbook, datos, "ventas_mes", "Ventas por mes")
    hoja_por_mes(workbook, datos, "compras_mes", "Compras por mes")
    hoja_resultados(workbook, datos)
    hoja_sumas_saldos(workbook, datos)
    hoja_libro(workbook, datos, "ventas", "Libro de ventas", "Cliente")
    hoja_libro(workbook, datos, "compras", "Libro de compras", "Proveedor")
    hoja_iva(workbook, datos)
    hoja_aging(workbook, datos, "por_cobrar", "Cuentas por cobrar", args.hasta, 1)
    hoja_aging(workbook, datos, "por_pagar", "Cuentas por pagar", args.hasta, -1)
    hoja_pagos(workbook, datos)

    destino.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destino)
    print(f"Listo: {destino}")
    print(f"  {len(datos['ventas'])} comprobantes de venta, {len(datos['compras'])} de compra")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
