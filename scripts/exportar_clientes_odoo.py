"""
Exporta los clientes de Odoo (res.partner) a un archivo Excel.

Columnas: Nombre, Razon social, Tipo de documento, RUT, Documento, Telefono,
Celular, Email e ID de Odoo.

Usa la configuracion ya guardada por la app (odoo_config.json), la misma que
emplea la integracion de expedicion.

Ejemplos:
    python scripts/exportar_clientes_odoo.py
    python scripts/exportar_clientes_odoo.py --salida "C:/Temp/clientes.xlsx"
    python scripts/exportar_clientes_odoo.py --solo-empresas --solo-con-rut
    python scripts/exportar_clientes_odoo.py --todos-los-contactos
"""

from __future__ import annotations

import argparse
import sys
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT / "config"))

from odoo_integration import odoo_integration  # noqa: E402

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - dependencia opcional
    print("Falta openpyxl. Instalalo con: pip install openpyxl")
    raise SystemExit(1)

# Tipos de identificacion uruguayos que representan un RUT de empresa.
RUT_TYPE_NAMES = {"RUT", "RUC", "RUT / RUC"}

PARTNER_FIELDS = [
    "id",
    "name",
    "commercial_company_name",
    "parent_id",
    "is_company",
    "vat",
    "l10n_latam_identification_type_id",
    "phone",
    "mobile",
    "email",
]

COLUMNS = [
    ("Nombre", 34),
    ("Razon social", 34),
    ("Tipo de documento", 18),
    ("RUT", 16),
    ("Documento", 16),
    ("Telefono", 18),
    ("Celular", 18),
    ("Email", 30),
    ("ID Odoo", 10),
]

BATCH_SIZE = 500


def _text(value: Any) -> str:
    """Odoo devuelve False para los campos char vacios."""
    if value is False or value is None:
        return ""
    return str(value).strip()


def _many2one_name(value: Any) -> str:
    if isinstance(value, (list, tuple)) and len(value) >= 2:
        return _text(value[1])
    return ""


def build_domain(args: argparse.Namespace) -> List[Any]:
    domain: List[Any] = []
    if not args.todos_los_contactos:
        domain.append(("customer_rank", ">", 0))
    if args.solo_empresas:
        domain.append(("is_company", "=", True))
    if args.solo_con_rut:
        domain.append(("vat", "!=", False))
    if args.incluir_archivados:
        domain.append(("active", "in", [True, False]))
    return domain


def fetch_partners(domain: List[Any], limite: int = 0) -> List[Dict[str, Any]]:
    uid = odoo_integration._authenticate()
    models = odoo_integration._xmlrpc_models()
    db = odoo_integration.config.get("database", "")
    password = odoo_integration.config.get("password", "")

    total = models.execute_kw(db, uid, password, "res.partner", "search_count", [domain])
    if limite:
        total = min(total, limite)
    print(f"Contactos a exportar: {total}")

    registros: List[Dict[str, Any]] = []
    offset = 0
    while offset < total:
        limit = min(BATCH_SIZE, total - offset)
        lote = models.execute_kw(
            db,
            uid,
            password,
            "res.partner",
            "search_read",
            [domain],
            {
                "fields": PARTNER_FIELDS,
                "offset": offset,
                "limit": limit,
                "order": "name asc, id asc",
            },
        )
        if not lote:
            break
        registros.extend(lote)
        offset += len(lote)
        print(f"  descargados {len(registros)}/{total}", end="\r", flush=True)
    print(f"  descargados {len(registros)}/{total}   ")
    return registros


def partner_to_row(partner: Dict[str, Any]) -> List[str]:
    nombre = _text(partner.get("name"))

    tipo_doc = _many2one_name(partner.get("l10n_latam_identification_type_id"))
    documento = _text(partner.get("vat"))
    es_rut = tipo_doc.upper() in RUT_TYPE_NAMES

    razon_social = _text(partner.get("commercial_company_name"))
    if not razon_social:
        # Un contacto con RUT es una entidad juridica aunque no tenga is_company.
        if partner.get("is_company") or es_rut:
            razon_social = nombre
        else:
            razon_social = _many2one_name(partner.get("parent_id"))

    telefono = _text(partner.get("phone"))
    celular = _text(partner.get("mobile"))
    if not telefono:
        telefono = celular

    return [
        nombre,
        razon_social,
        tipo_doc,
        documento if es_rut else "",
        documento,
        telefono,
        celular,
        _text(partner.get("email")),
        str(partner.get("id") or ""),
    ]


def escribir_excel(registros: List[Dict[str, Any]], destino: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Clientes"

    encabezado_fill = PatternFill("solid", fgColor="1F4E78")
    encabezado_font = Font(color="FFFFFF", bold=True)
    sheet.append([titulo for titulo, _ in COLUMNS])
    for celda in sheet[1]:
        celda.fill = encabezado_fill
        celda.font = encabezado_font
        celda.alignment = Alignment(vertical="center")

    for partner in registros:
        sheet.append(partner_to_row(partner))

    # RUT, documento y telefonos como texto: conservan ceros a la izquierda.
    for indice in (4, 5, 6, 7, 9):
        letra = get_column_letter(indice)
        for celda in sheet[letra][1:]:
            celda.number_format = "@"

    for indice, (_, ancho) in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(indice)].width = ancho

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{sheet.max_row}"

    destino.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destino)


def main() -> int:
    parser = argparse.ArgumentParser(description="Exporta clientes de Odoo a Excel")
    parser.add_argument("--salida", default="", help="Ruta del .xlsx a generar")
    parser.add_argument(
        "--todos-los-contactos",
        action="store_true",
        help="Exporta todos los contactos, no solo los que son clientes",
    )
    parser.add_argument("--solo-empresas", action="store_true", help="Solo contactos de tipo empresa")
    parser.add_argument("--solo-con-rut", action="store_true", help="Solo contactos con documento cargado")
    parser.add_argument("--incluir-archivados", action="store_true", help="Incluye contactos archivados")
    parser.add_argument("--limite", type=int, default=0, help="Maximo de contactos a exportar (0 = sin limite)")
    args = parser.parse_args()

    if not odoo_integration.is_configured():
        print("Odoo no esta configurado. Completa URL, base de datos, usuario y clave en la app.")
        return 1

    if args.salida:
        destino = Path(args.salida).expanduser()
    else:
        marca = datetime.now().strftime("%Y%m%d_%H%M")
        destino = Path.home() / "Desktop" / f"clientes_odoo_{marca}.xlsx"

    try:
        registros = fetch_partners(build_domain(args), limite=args.limite)
    except Exception as exc:
        print(f"Error consultando Odoo: {odoo_integration.humanize_exception(exc)}")
        return 1

    if not registros:
        print("No se encontraron contactos con esos filtros.")
        return 1

    escribir_excel(registros, destino)
    print(f"Listo: {len(registros)} clientes exportados a {destino}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
